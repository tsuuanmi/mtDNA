"""Tests for src.core.comparison — pairwise sample/batch comparison.

Covers concordance logic, unique-variant computation, summary/legacy row
formatting, variant-level rows, file writers, input validation, and the
end-to-end ``compare_batch_files`` entry point.
"""

import json
from pathlib import Path

import openpyxl
import pytest

from src.core.comparison import (
    ALWAYS_CONCORDANT_POSITIONS,
    RANGE_FLAGS,
    UNKNOWN_TOOL_NAME,
    SampleComparator,
    _batch_tool_name,
    _format_flag_reasons,
    _format_legacy_flagged,
    _format_tool_columns,
    _format_variant_display,
    _format_variant_flags,
    _get_position_flags,
    _legacy_summary_fields,
    _legacy_summary_row,
    _summary_fields,
    _tool_display_name,
    _tool_range_error,
    _validate_input_file,
    _variant_fields,
    compare_batch_files,
    create_variant_level_data,
    write_comparison_excel,
    write_comparison_json,
    write_comparison_tsv,
    write_variant_level_tsv,
)
from src.core.models import Sample, Tool, ToolResult, Variant
from src.core.variants import normalize_position

INT_FULL = {"HV1": [[16024, 16365]], "HV2": [[73, 340]], "HV3": [[438, 576]]}
INT_HV1_HV2 = {"HV1": [[16024, 16365]], "HV2": [[73, 340]]}
INT_GAP = {"HV2": [[73, 302], [316, 340]]}


# ---------------------------------------------------------------------------
# Helpers
# ---------------------------------------------------------------------------


def V(pos, ref, seq):
    """Build a normalized Variant."""
    return Variant(pos=normalize_position(pos), ref=ref, seq=seq)


def make_sample(
    sid,
    specs,
    *,
    tool=Tool.TRACY,
    intervals=INT_FULL,
    sflags=None,
    vflags=None,
):
    """Build a Sample from (pos, ref, seq) tuples."""
    return Sample(
        sample_id=sid,
        variants=[V(p, r, s) for p, r, s in specs],
        source_tool=tool,
        intervals=intervals,
        sample_flags=list(sflags or []),
        variant_flags=dict(vflags or {}),
        hv1=None,
        hv2=None,
        hv3=None,
        no_snps=None,
        no_ins=None,
        no_dels=None,
        no_identicals=None,
        no_unread=None,
        batch_id=None,
        information=None,
    )


def make_result(a, b):
    """Build a ComparisonResult via SampleComparator.compare."""
    return SampleComparator.compare(a, b)


# ---------------------------------------------------------------------------
# A. Tool display-name resolution
# ---------------------------------------------------------------------------


class TestToolDisplayName:
    def test_real_tool_uses_value(self):
        assert _tool_display_name(Tool.SEQUENCHER) == "sequencher"
        assert _tool_display_name(Tool.TRACY) == "tracy"

    def test_unknown_sentinel(self):
        assert _tool_display_name(Tool.UNKNOWN) == UNKNOWN_TOOL_NAME

    def test_none(self):
        assert _tool_display_name(None) == UNKNOWN_TOOL_NAME


class TestBatchToolName:
    def test_empty_batch_is_unknown(self):
        assert _batch_tool_name([]) == UNKNOWN_TOOL_NAME

    def test_first_sample_tool_used(self):
        batch = [make_sample("S1", [(73, "A", "G")], tool=Tool.SEQUENCHER)]
        assert _batch_tool_name(batch) == "sequencher"

    def test_unknown_tool_batch(self):
        batch = [make_sample("S1", [(73, "A", "G")], tool=Tool.UNKNOWN)]
        assert _batch_tool_name(batch) == UNKNOWN_TOOL_NAME


# ---------------------------------------------------------------------------
# B. Field schema column order
# ---------------------------------------------------------------------------


class TestFieldSchemas:
    def test_summary_fields_parameterized(self):
        f = _summary_fields("tracy", "sequencher")
        assert f[0] == "Sample ID"
        assert "Analyzed Range (tracy)" in f
        assert "Analyzed Range (sequencher)" in f
        assert "Concordant" in f
        assert "Flag" in f
        assert f.count("Sample ID") == 1
        # Count columns exist for both tools
        assert "Variants Count (tracy)" in f
        assert "Variants Count (sequencher)" in f

    def test_variant_fields_parameterized(self):
        f = _variant_fields("tracy", "sequencher")
        assert "tracy Variant" in f
        assert "sequencher Variant" in f
        assert "Region" in f
        assert "Position" in f

    def test_legacy_summary_fields_static_pipeline_sequencher(self):
        f = _legacy_summary_fields()
        assert f == [
            "Sample ID",
            "Batch",
            "Analyzed Range (Pipeline)",
            "Variants (Pipeline)",
            "Variants Unique (Pipeline)",
            "Analyzed Range (Sequencher)",
            "Variants (Sequencher)",
            "Variants Unique (Sequencher)",
            "Variants Count (Sequencher)",
            "Concordant",
            "Flag",
            "Flagged (Pipeline)",
            "Flagged (Sequencher)",
        ]
        # Legacy layout has a single Variants Count (Sequencher only).
        assert sum(1 for c in f if c.startswith("Variants Count")) == 1


# ---------------------------------------------------------------------------
# C. Range-error display (manual tools only)
# ---------------------------------------------------------------------------


class TestToolRangeError:
    def test_manual_tool_with_range_flag(self):
        tr = ToolResult(
            tool=Tool.MUTATION_SURVEYOR, analyzed_range="", variants=[], flag_reasons=["Range - Possibly wrong"]
        )
        assert _tool_range_error(tr) is True

    def test_manual_tool_missing_range_flag(self):
        tr = ToolResult(tool=Tool.SEQUENCHER, analyzed_range="", variants=[], flag_reasons=["Range - Missing"])
        assert _tool_range_error(tr) is True

    def test_manual_tool_other_flag_no_error(self):
        tr = ToolResult(tool=Tool.MUTATION_SURVEYOR, analyzed_range="", variants=[], flag_reasons=["Some other flag"])
        assert _tool_range_error(tr) is False

    def test_automated_tool_range_flag_ignored(self):
        # Pipeline/automated tools never get the Error display.
        tr = ToolResult(tool=Tool.TRACY, analyzed_range="", variants=[], flag_reasons=["Range - Possibly wrong"])
        assert _tool_range_error(tr) is False

    def test_range_flags_set_contents(self):
        assert "Range - Possibly wrong" in RANGE_FLAGS
        assert "Range - Missing" in RANGE_FLAGS


class TestFormatToolColumns:
    def test_range_error_shows_error(self):
        tr = ToolResult(
            tool=Tool.SEQUENCHER, analyzed_range="", variants=[{"pos": 73, "ref": "A", "seq": "G"}], flag_reasons=[]
        )
        cols = _format_tool_columns(tr, [{"pos": 73, "ref": "A", "seq": "G"}], range_error=True)
        assert cols == {"range": "Error", "variants": "Error", "unique": "Error", "count": 0}

    def test_normal_display(self):
        tr = ToolResult(
            tool=Tool.TRACY,
            analyzed_range="FULL REGION",
            variants=[{"pos": 73, "ref": "A", "seq": "G"}],
            flag_reasons=[],
        )
        cols = _format_tool_columns(tr, [{"pos": 73, "ref": "A", "seq": "G"}], range_error=False)
        assert cols["range"] == "FULL REGION"
        assert cols["variants"] == "73G"
        assert cols["unique"] == "73G"
        assert cols["count"] == 1

    def test_normal_display_empty_unique(self):
        tr = ToolResult(tool=Tool.TRACY, analyzed_range="None", variants=[], flag_reasons=[])
        cols = _format_tool_columns(tr, [], range_error=False)
        assert cols["unique"] == ""
        assert cols["count"] == 0


# ---------------------------------------------------------------------------
# D. Variant display + flag formatting helpers
# ---------------------------------------------------------------------------


class TestFormatVariantDisplay:
    def test_insertion_shows_seq(self):
        assert _format_variant_display({"pos": "309.1", "ref": "-", "seq": "C"}) == "C"

    def test_deletion_shows_del(self):
        assert _format_variant_display({"pos": 200, "ref": "A", "seq": "-"}) == "DEL"

    def test_snp_shows_alt(self):
        assert _format_variant_display({"pos": 73, "ref": "A", "seq": "G"}) == "G"

    def test_none_variant(self):
        assert _format_variant_display(None) == "None"

    def test_alt_alias_fallback(self):
        # Uses 'alt' when 'seq' absent.
        assert _format_variant_display({"pos": 73, "ref": "A", "alt": "T"}) == "T"


class TestGetPositionFlags:
    def test_lookup_by_position(self):
        vf = {"309.1|-|C": ["Insertion C at 309.1"], "310|A|G": ["Has 310 variant"]}
        assert _get_position_flags("309.1", None, vf) == "Insertion C at 309.1"

    def test_multiple_keys_same_position(self):
        vf = {"310|A|G": ["Has 310 variant"], "310|A|C": ["Heteroplasmy at 310"]}
        assert _get_position_flags("310", None, vf) == "Has 310 variant; Heteroplasmy at 310"

    def test_no_match(self):
        assert _get_position_flags("999", None, {"310|A|G": ["x"]}) == ""

    def test_empty_flags(self):
        assert _get_position_flags("310", None, {}) == ""


class TestFormatVariantFlags:
    def test_empty(self):
        assert _format_variant_flags({}) == ""

    def test_region_consolidation(self):
        vf = {
            "16185|A|G": ["16180-16193 region", "Heteroplasmy at 16185"],
            "16190|C|T": ["16180-16193 region"],
        }
        out = _format_variant_flags(vf)
        assert "Heteroplasmy at 16185" in out
        assert "16180-16193 region (16185, 16190)" in out
        # The bare per-variant region reason is consolidated, not duplicated.
        assert out.count("16180-16193 region") == 1

    def test_non_region_only(self):
        vf = {"73|A|G": ["Has 73 variant"]}
        assert _format_variant_flags(vf) == "Has 73 variant"

    def test_region_dedup_positions(self):
        # Same position appearing twice collapses to one position in the list.
        vf = {"16185|A|G": ["16180-16193 region"], "16185|A|C": ["16180-16193 region"]}
        assert _format_variant_flags(vf) == "16180-16193 region (16185)"


class TestFormatFlagReasons:
    def test_flagged_with_reasons(self):
        assert _format_flag_reasons(flagged=True, reasons=["a", "b"]) == "a; b"

    def test_flagged_no_reasons(self):
        assert _format_flag_reasons(flagged=True, reasons=[]) == "Yes"

    def test_not_flagged(self):
        assert _format_flag_reasons(flagged=False, reasons=["a"]) == "No"


class TestFormatLegacyFlagged:
    def test_merges_variant_then_sample(self):
        vf = {"309.1|-|C": ["Insertion C at 309.1"]}
        out = _format_legacy_flagged(["No 315.1 variant"], vf)
        # Variant-style flags come first.
        assert out == "Insertion C at 309.1, No 315.1 variant"

    def test_dedup(self):
        vf = {"73|A|G": ["Has 73 variant"]}
        out = _format_legacy_flagged(["Has 73 variant", "No 315.1 variant"], vf)
        # "Has 73 variant" appears once.
        assert out.count("Has 73 variant") == 1
        assert "No 315.1 variant" in out

    def test_empty_is_no(self):
        assert _format_legacy_flagged([], {}) == "No"


# ---------------------------------------------------------------------------
# E. format_intervals
# ---------------------------------------------------------------------------


class TestFormatIntervals:
    def test_full_region(self):
        assert SampleComparator.format_intervals(INT_FULL) == "FULL REGION"

    def test_partial_region(self):
        assert SampleComparator.format_intervals(INT_HV1_HV2) == "73-340 16024-16365"

    def test_none(self):
        assert SampleComparator.format_intervals(None) == "None"

    def test_empty_dict(self):
        assert SampleComparator.format_intervals({}) == "None"

    def test_empty_region_lists(self):
        assert SampleComparator.format_intervals({"HV2": []}) == "None"

    def test_multi_span_region(self):
        assert SampleComparator.format_intervals(INT_GAP) == "73-302 316-340"

    def test_reordered_ranges_sorted(self):
        out = SampleComparator.format_intervals({"HV2": [[316, 340], [73, 302]]})
        assert out == "73-302 316-340"


# ---------------------------------------------------------------------------
# F. find_unique_variants
# ---------------------------------------------------------------------------


class TestFindUniqueVariants:
    def test_empty_a(self):
        assert SampleComparator.find_unique_variants([], [V(73, "A", "G")]) == []

    def test_empty_b_returns_all_a(self):
        a = [V(73, "A", "G"), V(16223, "C", "T")]
        assert SampleComparator.find_unique_variants(a, []) == a

    def test_empty_b_ignore_special_filters_specials(self):
        a = [V(73, "A", "G"), V("309.1", "-", "C")]
        out = SampleComparator.find_unique_variants(a, [], ignore_special=True)
        assert out == [V(73, "A", "G")]

    def test_unique_in_a(self):
        a = [V(73, "A", "G"), V(16223, "C", "T")]
        b = [V(73, "A", "G")]
        assert SampleComparator.find_unique_variants(a, b) == [V(16223, "C", "T")]

    def test_ignore_special_skips_specials(self):
        a = [V(73, "A", "G"), V("309.1", "-", "C")]
        b = [V(73, "A", "G")]
        out = SampleComparator.find_unique_variants(a, b, ignore_special=True)
        assert out == []

    def test_position_normalization_match(self):
        # Integer 73 and string "73" normalize to the same key.
        a = [V(73, "A", "G")]
        b = [Variant(pos="73", ref="A", seq="G")]
        assert SampleComparator.find_unique_variants(a, b) == []

    def test_ref_seq_difference_means_not_shared(self):
        a = [V(73, "A", "G")]
        b = [V(73, "A", "T")]
        assert SampleComparator.find_unique_variants(a, b) == [V(73, "A", "G")]


# ---------------------------------------------------------------------------
# G. _non_special_variants + _collect_intervals
# ---------------------------------------------------------------------------


class TestNonSpecialVariants:
    def test_filters_special_positions(self):
        s = make_sample("S1", [(73, "A", "G"), ("309.1", "-", "C"), ("573.1", "-", "T")])
        out = SampleComparator._non_special_variants(s)
        assert [v.pos for v in out] == [73]

    def test_keeps_non_special(self):
        s = make_sample("S1", [(73, "A", "G"), (16223, "C", "T")])
        assert len(SampleComparator._non_special_variants(s)) == 2


class TestCollectIntervals:
    def test_collects_all_spans(self):
        s = make_sample("S1", [(73, "A", "G")], intervals=INT_FULL)
        assert sorted(SampleComparator._collect_intervals(s)) == sorted([(16024, 16365), (73, 340), (438, 576)])

    def test_multi_span(self):
        s = make_sample("S1", [(73, "A", "G")], intervals=INT_GAP)
        assert SampleComparator._collect_intervals(s) == [(73, 302), (316, 340)]

    def test_no_intervals_returns_empty(self):
        s = make_sample("S1", [(73, "A", "G")], intervals=None)
        assert SampleComparator._collect_intervals(s) == []


# ---------------------------------------------------------------------------
# H. _shared_positions_differ
# ---------------------------------------------------------------------------


class TestSharedPositionsDiffer:
    def test_identical_no_diff(self):
        a = [V(73, "A", "G"), V(16223, "C", "T")]
        b = [V(73, "A", "G"), V(16223, "C", "T")]
        assert SampleComparator._shared_positions_differ(a, b) is False

    def test_shared_differ(self):
        a = [V(73, "A", "G")]
        b = [V(73, "A", "T")]
        assert SampleComparator._shared_positions_differ(a, b) is True

    def test_always_concordant_ignored(self):
        # Differing ref/seq at an ALWAYS_CONCORDANT position is not a difference.
        a = [V("309.1", "-", "C")]
        b = [V("309.1", "-", "T")]
        assert SampleComparator._shared_positions_differ(a, b) is False
        assert "309.1" in ALWAYS_CONCORDANT_POSITIONS

    def test_non_shared_no_diff(self):
        a = [V(73, "A", "G")]
        b = [V(16223, "C", "T")]
        assert SampleComparator._shared_positions_differ(a, b) is False

    def test_ref_difference_counts(self):
        a = [V(73, "A", "G")]
        b = [V(73, "C", "G")]
        assert SampleComparator._shared_positions_differ(a, b) is True


# ---------------------------------------------------------------------------
# I. is_concordant
# ---------------------------------------------------------------------------


class TestIsConcordant:
    def test_identical_is_y(self):
        a = make_sample("S1", [(73, "A", "G"), (16223, "C", "T")])
        b = make_sample("S1", [(73, "A", "G"), (16223, "C", "T")])
        assert SampleComparator.is_concordant(a, b) == "Y"

    def test_shared_position_differ_is_n(self):
        a = make_sample("S1", [(73, "A", "G")])
        b = make_sample("S1", [(73, "A", "T")])
        assert SampleComparator.is_concordant(a, b) == "N"

    def test_always_concordant_differ_is_y(self):
        a = make_sample("S1", [("309.1", "-", "C")])
        b = make_sample("S1", [("309.1", "-", "T")])
        assert SampleComparator.is_concordant(a, b) == "Y"

    def test_a_only_in_hv_is_n(self):
        a = make_sample("S1", [(73, "A", "G"), (16223, "C", "T")])
        b = make_sample("S1", [(73, "A", "G")])
        assert SampleComparator.is_concordant(a, b) == "N"

    def test_a_only_outside_hv_is_y(self):
        a = make_sample("S1", [(73, "A", "G"), (5000, "A", "G")])
        b = make_sample("S1", [(73, "A", "G")])
        assert SampleComparator.is_concordant(a, b) == "Y"

    def test_b_only_in_hv_within_intervals_is_n(self):
        a = make_sample("S1", [(73, "A", "G")])
        b = make_sample("S1", [(73, "A", "G"), (16223, "C", "T")])
        assert SampleComparator.is_concordant(a, b) == "N"

    def test_b_only_outside_hv_is_y(self):
        a = make_sample("S1", [(73, "A", "G")])
        b = make_sample("S1", [(73, "A", "G"), (5000, "A", "G")])
        assert SampleComparator.is_concordant(a, b) == "Y"

    def test_b_only_in_hv_no_a_intervals_is_n(self):
        # When A has no intervals, B-only HV variants are treated as in-range -> N.
        a = make_sample("S1", [(73, "A", "G")], intervals=None)
        b = make_sample("S1", [(73, "A", "G"), (16223, "C", "T")])
        assert SampleComparator.is_concordant(a, b) == "N"

    def test_b_only_outside_a_intervals_is_y(self):
        # B-only HV variant falls in a gap of A's analyzed ranges -> not flagged.
        a = make_sample("S1", [(73, "A", "G")], intervals=INT_GAP)
        b = make_sample("S1", [(73, "A", "G"), (310, "T", "C")])
        # 310 is in HV2 (73-340) but in A's gap (303-315 missing) -> Y.
        assert SampleComparator.is_concordant(a, b) == "Y"

    def test_b_only_insertion_after_a_interval_end_is_y(self):
        a = make_sample("S1", [(16189, "T", "C")], intervals={"HV1": [[16024, 16193]]})
        b = make_sample("S1", [(16189, "T", "C"), ("16193.1", "-", "C")])
        assert SampleComparator.is_concordant(a, b) == "Y"

    def test_empty_a_is_na(self):
        a = make_sample("S1", [])
        b = make_sample("S1", [(73, "A", "G")])
        assert SampleComparator.is_concordant(a, b) == "N/A"

    def test_empty_b_is_na(self):
        a = make_sample("S1", [(73, "A", "G")])
        b = make_sample("S1", [])
        assert SampleComparator.is_concordant(a, b) == "N/A"

    def test_special_positions_excluded_from_concordance(self):
        # Special-position insertions do not affect concordance.
        a = make_sample("S1", [("309.1", "-", "C"), (73, "A", "G")])
        b = make_sample("S1", [("309.1", "-", "C"), (73, "A", "G")])
        assert SampleComparator.is_concordant(a, b) == "Y"


# ---------------------------------------------------------------------------
# J. compare() full ComparisonResult
# ---------------------------------------------------------------------------


class TestCompare:
    def test_concordance_propagated(self):
        a = make_sample("S1", [(73, "A", "G")], sflags=["No 315.1 variant"])
        b = make_sample("S1", [(73, "A", "G")], tool=Tool.SEQUENCHER)
        r = make_result(a, b)
        assert r.concordance == "Y"
        assert r.sample_id == "S1"

    def test_flagged_from_sample_or_variant_flags(self):
        # A tool is flagged when it has sample flags or variant flags.
        a = make_sample("S1", [(73, "A", "G")], sflags=["No 315.1 variant"])
        b = make_sample("S1", [(73, "A", "G")], tool=Tool.SEQUENCHER, vflags={"73|A|G": ["Has 73 variant"]})
        r = make_result(a, b)
        assert r.tool_a.flagged is True
        assert r.tool_b.flagged is True

    def test_not_flagged_when_clean(self):
        a = make_sample("S1", [(73, "A", "G")])
        b = make_sample("S1", [(73, "A", "G")], tool=Tool.SEQUENCHER)
        r = make_result(a, b)
        assert r.tool_a.flagged is False
        assert r.tool_b.flagged is False

    def test_unique_variants_computed(self):
        a = make_sample("S1", [(73, "A", "G"), (5000, "A", "G")])
        b = make_sample("S1", [(73, "A", "G")], tool=Tool.SEQUENCHER)
        r = make_result(a, b)
        assert [v["pos"] for v in r.unique_variants_a] == [5000]
        assert r.unique_variants_b == []

    def test_tool_variants_are_dicts(self):
        a = make_sample("S1", [(73, "A", "G")])
        b = make_sample("S1", [(73, "A", "G")], tool=Tool.SEQUENCHER)
        r = make_result(a, b)
        assert r.tool_a.variants[0]["pos"] == 73
        assert r.tool_a.variants[0]["ref"] == "A"
        assert r.tool_b.variants[0]["seq"] == "G"

    def test_analyzed_range_propagated(self):
        a = make_sample("S1", [(73, "A", "G")])
        b = make_sample("S1", [(73, "A", "G")], tool=Tool.SEQUENCHER)
        r = make_result(a, b)
        assert r.tool_a.analyzed_range == "FULL REGION"
        assert r.tool_b.analyzed_range == "FULL REGION"

    def test_flag_reasons_copied(self):
        a = make_sample("S1", [(73, "A", "G")], sflags=["No 315.1 variant", "Has 309T variant"])
        b = make_sample("S1", [(73, "A", "G")], tool=Tool.SEQUENCHER)
        r = make_result(a, b)
        assert r.tool_a.flag_reasons == ["No 315.1 variant", "Has 309T variant"]
        assert r.tool_b.flag_reasons == []


# ---------------------------------------------------------------------------
# K. summary_row
# ---------------------------------------------------------------------------


class TestSummaryRow:
    def test_row_keys_and_flag(self):
        a = make_sample("S1", [(73, "A", "G")], sflags=["No 315.1 variant"])
        b = make_sample("S1", [(73, "A", "G")], tool=Tool.SEQUENCHER)
        r = make_result(a, b)
        row = SampleComparator.summary_row(r, batch_id="B1", tool_name_a="tracy", tool_name_b="sequencher")
        assert row["Sample ID"] == "S1"
        assert row["Batch"] == "B1"
        assert row["Concordant"] == "Y"
        assert row["Flag"] == "Yes"
        assert row["Analyzed Range (tracy)"] == "FULL REGION"
        assert row["Variants (tracy)"] == "73G"
        assert row["Variants Count (tracy)"] == 1

    def test_flag_no_when_clean(self):
        a = make_sample("S1", [(73, "A", "G")])
        b = make_sample("S1", [(73, "A", "G")], tool=Tool.SEQUENCHER)
        r = make_result(a, b)
        row = SampleComparator.summary_row(r, tool_name_a="tracy", tool_name_b="sequencher")
        assert row["Flag"] == "No"
        assert row["Sample Flags (tracy)"] == "No"

    def test_range_error_displayed(self):
        a = make_sample("S1", [(73, "A", "G")], tool=Tool.SEQUENCHER, sflags=["Range - Possibly wrong"])
        b = make_sample("S1", [(73, "A", "G")])
        r = make_result(a, b)
        row = SampleComparator.summary_row(r, tool_name_a="sequencher", tool_name_b="tracy")
        assert row["Analyzed Range (sequencher)"] == "Error"
        assert row["Variants (sequencher)"] == "Error"
        assert row["Variants Unique (sequencher)"] == "Error"
        # Automated tool (tracy) columns unaffected.
        assert row["Analyzed Range (tracy)"] == "FULL REGION"

    def test_sample_flag_reasons_formatted(self):
        a = make_sample("S1", [(73, "A", "G")], sflags=["No 315.1 variant", "Has 309T variant"])
        b = make_sample("S1", [(73, "A", "G")], tool=Tool.SEQUENCHER)
        r = make_result(a, b)
        row = SampleComparator.summary_row(r, tool_name_a="tracy", tool_name_b="sequencher")
        assert row["Sample Flags (tracy)"] == "No 315.1 variant; Has 309T variant"


# ---------------------------------------------------------------------------
# L. legacy summary row
# ---------------------------------------------------------------------------


class TestLegacySummaryRow:
    def test_row_has_13_columns(self):
        a = make_sample("S1", [(73, "A", "G")], sflags=["No 315.1 variant"])
        b = make_sample("S1", [(73, "A", "G")], tool=Tool.SEQUENCHER)
        r = make_result(a, b)
        row = _legacy_summary_row(r, batch_id="B1")
        assert set(row.keys()) == set(_legacy_summary_fields())

    def test_pipeline_sequencher_labels(self):
        a = make_sample("S1", [(73, "A", "G")])
        b = make_sample("S1", [(73, "A", "G")], tool=Tool.SEQUENCHER)
        r = make_result(a, b)
        row = _legacy_summary_row(r, batch_id="B1")
        assert "Analyzed Range (Pipeline)" in row
        assert "Analyzed Range (Sequencher)" in row
        assert "Variants Count (Sequencher)" in row
        assert row["Concordant"] == "Y"

    def test_flagged_columns_merge(self):
        a = make_sample(
            "S1",
            [(73, "A", "G")],
            sflags=["No 315.1 variant"],
            vflags={"73|A|G": ["Has 73 variant"]},
        )
        b = make_sample("S1", [(73, "A", "G")], tool=Tool.SEQUENCHER)
        r = make_result(a, b)
        row = _legacy_summary_row(r, batch_id="B1")
        # Variant flag first, then sample flag.
        assert row["Flagged (Pipeline)"] == "Has 73 variant, No 315.1 variant"
        assert row["Flagged (Sequencher)"] == "No"

    def test_flag_yes_when_either_flagged(self):
        a = make_sample("S1", [(73, "A", "G")], sflags=["No 315.1 variant"])
        b = make_sample("S1", [(73, "A", "G")], tool=Tool.SEQUENCHER)
        r = make_result(a, b)
        assert _legacy_summary_row(r)["Flag"] == "Yes"


# ---------------------------------------------------------------------------
# M. compare_batches
# ---------------------------------------------------------------------------


def _batch(*samples: Sample) -> list[Sample]:
    return list(samples)


class TestCompareBatches:
    def test_common_a_only_b_only(self):
        ba = _batch(
            make_sample("S1", [(73, "A", "G")]),
            make_sample("S2", [(16223, "C", "T")]),
        )
        bb = _batch(
            make_sample("S1", [(73, "A", "G")], tool=Tool.SEQUENCHER),
            make_sample("S3", [(490, "T", "C")], tool=Tool.SEQUENCHER),
        )
        res = SampleComparator.compare_batches(ba, bb, batch_id="B1")
        assert res["summary"]["matched"] == 1
        assert res["summary"]["a_only"] == ["S2"]
        assert res["summary"]["b_only"] == ["S3"]

    def test_summary_counts(self):
        ba = _batch(
            make_sample("S1", [(73, "A", "G")]),
            make_sample("S2", [(73, "A", "G")]),  # concordant
            make_sample("S3", [(73, "A", "G")]),
        )
        bb = _batch(
            make_sample("S1", [(73, "A", "G")], tool=Tool.SEQUENCHER, sflags=["No 315.1 variant"]),  # flagged
            make_sample("S2", [(73, "A", "G")], tool=Tool.SEQUENCHER),  # clean, concordant
            make_sample("S3", [(73, "A", "T")], tool=Tool.SEQUENCHER),  # discordant
        )
        res = SampleComparator.compare_batches(ba, bb, batch_id="B1")
        assert res["summary"]["matched"] == 3
        assert res["summary"]["concordant"] == 2  # S1, S2
        assert res["summary"]["discordant"] == 1  # S3
        assert res["summary"]["flagged"] == 1  # S1 (b has flags)

    def test_tool_names_from_batches(self):
        ba = _batch(make_sample("S1", [(73, "A", "G")], tool=Tool.TRACY))
        bb = _batch(make_sample("S1", [(73, "A", "G")], tool=Tool.SEQUENCHER))
        res = SampleComparator.compare_batches(ba, bb, batch_id="B1")
        assert res["tool_name_a"] == "tracy"
        assert res["tool_name_b"] == "sequencher"

    def test_empty_batch_unknown_tool_name(self):
        ba = []
        bb = _batch(make_sample("S1", [(73, "A", "G")], tool=Tool.SEQUENCHER))
        res = SampleComparator.compare_batches(ba, bb, batch_id="B1")
        assert res["tool_name_a"] == UNKNOWN_TOOL_NAME
        assert res["tool_name_b"] == "sequencher"
        assert res["summary"]["matched"] == 0
        assert res["summary"]["a_only"] == []
        assert res["summary"]["b_only"] == ["S1"]

    def test_unknown_tool_batch(self):
        ba = _batch(make_sample("S1", [(73, "A", "G")], tool=Tool.UNKNOWN))
        bb = _batch(make_sample("S1", [(73, "A", "G")], tool=Tool.UNKNOWN))
        res = SampleComparator.compare_batches(ba, bb, batch_id="B1")
        assert res["tool_name_a"] == UNKNOWN_TOOL_NAME
        assert res["tool_name_b"] == UNKNOWN_TOOL_NAME

    def test_legacy_format_fields(self):
        ba = _batch(make_sample("S1", [(73, "A", "G")]))
        bb = _batch(make_sample("S1", [(73, "A", "G")], tool=Tool.SEQUENCHER))
        res = SampleComparator.compare_batches(ba, bb, batch_id="B1", legacy_format=True)
        assert res["legacy_format"] is True
        assert res["summary_fields"] == _legacy_summary_fields()
        assert res["summary_rows"][0].keys() == set(_legacy_summary_fields())

    def test_default_format_fields(self):
        ba = _batch(make_sample("S1", [(73, "A", "G")], tool=Tool.TRACY))
        bb = _batch(make_sample("S1", [(73, "A", "G")], tool=Tool.SEQUENCHER))
        res = SampleComparator.compare_batches(ba, bb, batch_id="B1")
        assert res["summary_fields"] == _summary_fields("tracy", "sequencher")
        assert res["variant_fields"] == _variant_fields("tracy", "sequencher")

    def test_variant_rows_aggregated_across_samples(self):
        ba = _batch(
            make_sample("S1", [(73, "A", "G")]),
            make_sample("S2", [(16223, "C", "T")]),
        )
        bb = _batch(
            make_sample("S1", [(73, "A", "G")], tool=Tool.SEQUENCHER),
            make_sample("S2", [(16223, "C", "T")], tool=Tool.SEQUENCHER),
        )
        res = SampleComparator.compare_batches(ba, bb, batch_id="B1")
        # 2 samples x 1 variant each = 2 rows
        assert len(res["variant_rows"]) == 2
        assert res["variant_rows"][0]["Sample ID"] == "S1"


# ---------------------------------------------------------------------------
# N. create_variant_level_data
# ---------------------------------------------------------------------------


class TestCreateVariantLevelData:
    def test_union_of_positions(self):
        a = make_sample("S1", [(73, "A", "G"), (16223, "C", "T")])
        b = make_sample("S1", [(73, "A", "G"), (490, "T", "C")], tool=Tool.SEQUENCHER)
        r = make_result(a, b)
        rows = create_variant_level_data(r, batch_id="B1", tool_name_a="tracy", tool_name_b="sequencher")
        positions = [row["Position"] for row in rows]
        assert positions == ["73", "490", "16223"]
        assert rows[0]["Region"] == "HV2"
        assert rows[1]["Region"] == "HV3"
        assert rows[2]["Region"] == "HV1"

    def test_a_only_position_shows_none_for_b(self):
        a = make_sample("S1", [(73, "A", "G"), (5000, "A", "G")])
        b = make_sample("S1", [(73, "A", "G")], tool=Tool.SEQUENCHER)
        r = make_result(a, b)
        rows = {
            row["Position"]: row for row in create_variant_level_data(r, tool_name_a="tracy", tool_name_b="sequencher")
        }
        assert rows["5000"]["tracy Variant"] == "G"
        assert rows["5000"]["sequencher Variant"] == "None"
        assert rows["5000"]["Region"] == "Other"

    def test_variant_display_insertion_deletion_snp(self):
        a = make_sample("S1", [("309.1", "-", "C"), (200, "A", "-"), (73, "A", "G")])
        b = make_sample("S1", [("309.1", "-", "C"), (200, "A", "-"), (73, "A", "G")], tool=Tool.SEQUENCHER)
        r = make_result(a, b)
        rows = {
            row["Position"]: row for row in create_variant_level_data(r, tool_name_a="tracy", tool_name_b="sequencher")
        }
        # Note: 309.1 is a special position; it still appears in variant-level
        # display (create_variant_level_data does not filter specials).
        assert rows["309.1"]["tracy Variant"] == "C"
        assert rows["200"]["tracy Variant"] == "DEL"
        assert rows["73"]["tracy Variant"] == "G"

    def test_position_flags_attached(self):
        a = make_sample("S1", [(73, "A", "G")], vflags={"73|A|G": ["Has 73 variant"]})
        b = make_sample("S1", [(73, "A", "G")], tool=Tool.SEQUENCHER)
        r = make_result(a, b)
        rows = create_variant_level_data(r, tool_name_a="tracy", tool_name_b="sequencher")
        assert rows[0]["Variant Flags (tracy)"] == "Has 73 variant"
        assert rows[0]["Variant Flags (sequencher)"] == ""

    def test_sample_flags_in_rows(self):
        a = make_sample("S1", [(73, "A", "G")], sflags=["No 315.1 variant"])
        b = make_sample("S1", [(73, "A", "G")], tool=Tool.SEQUENCHER)
        r = make_result(a, b)
        rows = create_variant_level_data(r, tool_name_a="tracy", tool_name_b="sequencher")
        assert rows[0]["Sample Flags (tracy)"] == "No 315.1 variant"
        assert rows[0]["Sample Flags (sequencher)"] == ""


# ---------------------------------------------------------------------------
# O. File writers
# ---------------------------------------------------------------------------


@pytest.fixture
def batch_result():
    ba = _batch(
        make_sample("S1", [(73, "A", "G")], sflags=["No 315.1 variant"]),
        make_sample("S2", [(16223, "C", "T")]),
    )
    bb = _batch(
        make_sample("S1", [(73, "A", "G")], tool=Tool.SEQUENCHER),
        make_sample("S2", [(16223, "C", "T")], tool=Tool.SEQUENCHER),
    )
    return SampleComparator.compare_batches(ba, bb, batch_id="B1")


class TestWriters:
    def test_write_comparison_json(self, tmp_path, batch_result):
        out = write_comparison_json(batch_result, tmp_path / "out.json")
        data = json.loads(out.read_text(encoding="utf-8"))
        assert "summary" in data
        assert len(data["summary_rows"]) == 2
        assert data["summary"]["matched"] == 2

    def test_write_comparison_tsv(self, tmp_path, batch_result):
        out = write_comparison_tsv(batch_result, tmp_path / "out.tsv")
        lines = out.read_text(encoding="utf-8").splitlines()
        assert lines[0].split("\t")[0] == "Sample ID"
        assert len(lines) == 3  # header + 2 rows
        assert "Concordant" in lines[0]

    def test_write_variant_level_tsv(self, tmp_path, batch_result):
        out = write_variant_level_tsv(batch_result, tmp_path / "variants.tsv")
        lines = out.read_text(encoding="utf-8").splitlines()
        assert lines[0].split("\t")[0] == "Sample ID"
        assert "Region" in lines[0]
        assert "Position" in lines[0]

    def test_write_comparison_excel(self, tmp_path, batch_result):
        out = write_comparison_excel(batch_result, tmp_path / "out.xlsx")
        assert out.exists()
        # Read back to confirm structure.
        wb = openpyxl.load_workbook(out)
        ws = wb["Comparison Results"]
        assert ws.cell(1, 1).value == "Sample ID"
        assert ws.cell(2, 1).value in {"S1", "S2"}

    def test_write_comparison_excel_empty_rows(self, tmp_path):
        res = SampleComparator.compare_batches([], [], batch_id="B1")
        out = write_comparison_excel(res, tmp_path / "empty.xlsx")
        # Empty rows -> early return; the path is returned but no file is written.
        assert out == tmp_path / "empty.xlsx"
        assert not out.exists()


# ---------------------------------------------------------------------------
# P. _validate_input_file
# ---------------------------------------------------------------------------


class TestValidateInputFile:
    def test_valid_json(self, tmp_path):
        p = tmp_path / "a.json"
        p.write_text('{"S1": {}}', encoding="utf-8")
        assert _validate_input_file(p, "Batch A") is True

    def test_missing_file(self, tmp_path):
        assert _validate_input_file(tmp_path / "missing.json", "Batch A") is False

    def test_invalid_json(self, tmp_path):
        p = tmp_path / "bad.json"
        p.write_text("{not json", encoding="utf-8")
        assert _validate_input_file(p, "Batch A") is False


# ---------------------------------------------------------------------------
# Q. compare_batch_files end-to-end
# ---------------------------------------------------------------------------


def _write_batch_json(path: Path, samples_data: dict[str, object]) -> None:
    """Write a batch JSON file loadable by load_sample_batch.

    samples_data: dict {sample_id: {variants: [...], intervals: {...}, source_tool: str}}
    """
    path.write_text(json.dumps(samples_data), encoding="utf-8")


class TestCompareBatchFiles:
    def test_end_to_end_outputs(self, tmp_path):
        batch_a_path = tmp_path / "batch_a.json"
        batch_b_path = tmp_path / "batch_b.json"
        _write_batch_json(
            batch_a_path,
            {
                "S1": {
                    "variants": [{"pos": 73, "ref": "A", "seq": "G"}],
                    "intervals": INT_FULL,
                    "source_tool": "tracy",
                },
            },
        )
        _write_batch_json(
            batch_b_path,
            {
                "S1": {
                    "variants": [{"pos": 73, "ref": "A", "seq": "G"}],
                    "intervals": INT_FULL,
                    "source_tool": "sequencher",
                },
            },
        )
        out_dir = tmp_path / "out"
        res = compare_batch_files(batch_a_path, batch_b_path, out_dir, batch_id="B1")
        assert res["batch_id"] == "B1"
        assert (out_dir / "B1.json").exists()
        assert (out_dir / "B1.tsv").exists()
        assert (out_dir / "B1_variant_level.tsv").exists()
        assert (out_dir / "B1.xlsx").exists()
        assert res["summary"]["matched"] == 1
        assert res["summary"]["concordant"] == 1

    def test_legacy_format_output(self, tmp_path):
        batch_a_path = tmp_path / "batch_a.json"
        batch_b_path = tmp_path / "batch_b.json"
        _write_batch_json(
            batch_a_path,
            {
                "S1": {
                    "variants": [{"pos": 73, "ref": "A", "seq": "G"}],
                    "intervals": INT_FULL,
                    "source_tool": "tracy",
                },
            },
        )
        _write_batch_json(
            batch_b_path,
            {
                "S1": {
                    "variants": [{"pos": 73, "ref": "A", "seq": "G"}],
                    "intervals": INT_FULL,
                    "source_tool": "sequencher",
                },
            },
        )
        res = compare_batch_files(batch_a_path, batch_b_path, tmp_path / "out", batch_id="B1", legacy_format=True)
        assert res["legacy_format"] is True
        assert res["summary_fields"] == _legacy_summary_fields()
