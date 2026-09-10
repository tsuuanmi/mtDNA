#!/usr/bin/env python3
"""
ms_trace_process.py

MS trace-processing version based on MS_all_blocks_compare_QC_cli_v13.py.
Truthset-dependent logic has been removed.

Keeps:
1) Multi-block MS custom report parsing.
2) Cleaning to remove non-real trace rows.
3) Variant Summary / Variant Converted / Variant Converted 2.
4) QC-ready repeat-region flags:
   - HV1_PolyC_flag (HV1)
   - HV2-3_PolyC_flag (HV2-3)

Removes:
1) --truth input.
2) Truthset merge/splitting.
3) Concordance_auto / Concordance_auto_2.
4) Error_type, because it depends on concordance.
"""

import argparse
import os
import re

import pandas as pd
from openpyxl import load_workbook


# =========================
# INPUT / OUTPUT
# =========================
def parse_args():
    p = argparse.ArgumentParser(
        description="Process MS custom report into trace-level normalized MS output without truthset comparison"
    )
    p.add_argument("--ms", required=True, help="MS custom report xlsx")
    p.add_argument("--lid-manifest", required=True, help="Excel file containing expected LID list for this batch")
    p.add_argument("--inventory-out", default=None, help="Optional Layer 0 inventory QC report xlsx")
    p.add_argument("--out", required=True, help="Output MS trace processed xlsx, e.g. {BATCH}_MS_trace_processed.xlsx")
    return p.parse_args()


args = parse_args()
MS_FILE = args.ms
OUTPUT = args.out
LID_MANIFEST = args.lid_manifest
INVENTORY_OUTPUT = args.inventory_out


# =========================
# CONSTANTS
# =========================
IUPAC = {
    "W": set("AT"),
    "R": set("AG"),
    "M": set("AC"),
    "K": set("GT"),
    "Y": set("CT"),
    "S": set("CG"),
    "D": set("AGT"),
    "H": set("ACT"),
    "V": set("ACG"),
    "B": set("CGT"),
}

UNORDERED_MAP = {
    "".join(sorted("AT")): "W",
    "".join(sorted("AG")): "R",
    "".join(sorted("AC")): "M",
    "".join(sorted("GT")): "K",
    "".join(sorted("CT")): "Y",
    "".join(sorted("CG")): "S",
    "".join(sorted("AGT")): "D",
    "".join(sorted("ACT")): "H",
    "".join(sorted("ACG")): "V",
    "".join(sorted("CGT")): "B",
}


# =========================
# READ ALL HEADER BLOCKS
# =========================
def find_all_header_rows(ws, scan_rows=500):
    header_rows = []
    max_rows = min(scan_rows, ws.max_row)

    for r in range(1, max_rows + 1):
        vals = [ws.cell(r, c).value for c in range(1, ws.max_column + 1)]
        text = [str(v).strip() if v is not None else "" for v in vals]

        if "Trace #" in text and "Sample Name" in text:
            header_rows.append(r)

    if not header_rows:
        raise ValueError("Could not detect any header rows.")
    return header_rows


def header_score(ws, row):
    vals = [ws.cell(row, c).value for c in range(1, ws.max_column + 1)]
    text = [str(v).strip() if v is not None else "" for v in vals]
    return sum(1 for x in text if re.fullmatch(r"Variant\d+", x))


def get_headers(ws, row):
    headers = [ws.cell(row, c).value for c in range(1, ws.max_column + 1)]
    return [str(h).strip() if h is not None else f"Unnamed_{i + 1}" for i, h in enumerate(headers)]


def merge_headers(master_headers, block_headers):
    out = []
    max_len = max(len(master_headers), len(block_headers))
    for i in range(max_len):
        m = master_headers[i] if i < len(master_headers) else f"Unnamed_{i + 1}"
        b = block_headers[i] if i < len(block_headers) else ""
        out.append(b if b and b != f"Unnamed_{i + 1}" else m)
    return out


def read_ms_all_blocks(path):
    wb = load_workbook(path, data_only=True)
    ws = wb[wb.sheetnames[0]]

    header_rows = find_all_header_rows(ws)

    best_header = max(header_rows, key=lambda r: header_score(ws, r))
    master_headers = get_headers(ws, best_header)

    rows = []

    for idx, hr in enumerate(header_rows):
        next_hr = header_rows[idx + 1] if idx + 1 < len(header_rows) else ws.max_row + 1
        block_headers = get_headers(ws, hr)
        headers = merge_headers(master_headers, block_headers)

        try:
            # trace_idx = headers.index("Trace #") + 1
            sample_idx = headers.index("Sample Name") + 1
        except ValueError:
            continue

        for r in range(hr + 1, next_hr):
            sample = ws.cell(r, sample_idx).value

            if isinstance(sample, str) and ".ab1" in sample:
                row = [ws.cell(r, c).value for c in range(1, ws.max_column + 1)]
                if len(row) < len(headers):
                    row += [None] * (len(headers) - len(row))
                elif len(row) > len(headers):
                    row = row[: len(headers)]
                rows.append(row)

    if not rows:
        return pd.DataFrame(columns=master_headers)

    df = pd.DataFrame(rows, columns=master_headers[: len(rows[0])])
    df.columns = [str(c).strip() for c in df.columns]
    df = df.drop_duplicates()
    return df


# =========================
# QC LAYER 0 - INVENTORY CHECK
# =========================
EXPECTED_PRIMERS = ["HV1F", "HV1R", "HV2F", "HV3R"]


def is_control_lid(lid):
    """Return True for control LIDs that are expected in MS reports but not in the LID manifest.

    Supported formats:
      - NTC1, NTC2, ...
      - PC1_mtDNA_247547, PC2_mtDNA_247547, ...
    """
    s = str(lid).strip()
    return (
        re.fullmatch(r"NTC\d+", s, flags=re.IGNORECASE) is not None
        or re.fullmatch(r"PC\d+_.+", s, flags=re.IGNORECASE) is not None
    )


def _as_display_string(v):
    """Convert Excel cell value to a stable display string without stripping spaces."""
    if pd.isna(v):
        return ""
    if isinstance(v, float) and v.is_integer():
        return str(int(v))
    return str(v)


def _batch_from_manifest_filename(path):
    """Extract BATCH from manifest filename.

    Accepted patterns:
      - {BATCH_ID}.xlsx  (metadata directory, e.g. 20260504_mtDNA_316.xlsx)
      - YYYYMMDD_Danh sách mẫu_{BATCH}.xlsx  (e.g. 20260515_Danh sách mẫu_MS_210426_003.xlsx)
    """
    name = os.path.basename(path)
    # Date-prefixed pattern first (more specific)
    m = re.match(r"^\d{8}_Danh sách mẫu_(.+)\.xlsx$", name, flags=re.IGNORECASE)
    if m:
        return m.group(1)
    # Metadata directory pattern: {BATCH_ID}.xlsx
    m = re.match(r"^(.+)\.xlsx$", name, flags=re.IGNORECASE)
    if m:
        return m.group(1)
    return ""


def _manifest_filename_matches_required_pattern(path):
    name = os.path.basename(path)
    # Accept both: YYYYMMDD_Danh sách mẫu_{BATCH}.xlsx and {BATCH}.xlsx (metadata dir)
    return (
        re.match(r"^\d{8}_Danh sách mẫu_.+\.xlsx$", name, flags=re.IGNORECASE) is not None
        or re.match(r"^.+\.xlsx$", name, flags=re.IGNORECASE) is not None
    )


def _custom_report_filename_matches_batch(path, expected_batch):
    """
    Required custom report filename:
      {BATCH}_anything_df_custom_report.xlsx

    The middle `anything` segment may be empty, so these are valid:
      {BATCH}_df_custom_report.xlsx
      {BATCH}_rerun_df_custom_report.xlsx
    """
    name = os.path.basename(path)
    if not expected_batch:
        return False
    pattern = r"^" + re.escape(expected_batch) + r"(?:_.+)?_df_custom_report\.xlsx$"
    return re.match(pattern, name, flags=re.IGNORECASE) is not None


def _default_inventory_output(output_path):
    root, ext = os.path.splitext(output_path)
    if ext.lower() not in [".xlsx", ".xlsm"]:
        return output_path + "_inventory_QC.xlsx"
    return root + "_inventory_QC.xlsx"


def read_lid_manifest(path):
    df = pd.read_excel(path)
    df.columns = [str(c).strip() for c in df.columns]

    required = ["LID", "MẺ CHẠY"]
    missing_cols = [c for c in required if c not in df.columns]
    if missing_cols:
        raise ValueError(f"LID manifest missing required column(s): {missing_cols}")

    out = pd.DataFrame()
    out["LID_raw"] = df["LID"].apply(_as_display_string)
    out["Batch_raw"] = df["MẺ CHẠY"].apply(_as_display_string)
    out["Row_Number"] = df.index + 2  # Excel row number when header is row 1
    return out


def run_inventory_qc(ms_df, manifest_path, output_path, ms_path):
    """
    Layer 0 is a strict batch integrity gate.

    Any error below fails the whole batch and prevents MS_trace_processed generation:
    - manifest filename not matching {BATCH}.xlsx or YYYYMMDD_Danh sách mẫu_{BATCH}.xlsx
    - custom report filename not matching {BATCH}_anything_df_custom_report.xlsx
    - batch code mismatch between manifest filename and MẺ CHẠY column
    - blank LID / batch cells in manifest
    - leading/trailing spaces in manifest LID or batch code
    - duplicate LID in manifest
    - missing sample LID or extra non-control LID in MS custom report
    - PC/NTC control LIDs are allowed in MS report and excluded from manifest matching/count checks
    - blank/unparsed LID or primer in MS custom report
    - unexpected primer in MS custom report
    - missing or duplicate expected primer per LID
    - total primer count != number of manifest LIDs
    """
    inv_out = INVENTORY_OUTPUT or _default_inventory_output(output_path)
    manifest = read_lid_manifest(manifest_path)
    details = []

    def add_issue(issue_type, lid="", primer="", details_text="", severity="ERROR", row_number=""):
        details.append(
            {
                "Severity": severity,
                "Issue_Type": issue_type,
                "LID": lid,
                "Primer": primer,
                "Details": details_text,
                "Manifest_Row": row_number,
            }
        )

    # File identity checks. These are blocking Layer 0 errors.
    ms_filename = os.path.basename(ms_path)
    manifest_filename = os.path.basename(manifest_path)
    if not _manifest_filename_matches_required_pattern(manifest_path):
        add_issue(
            "Manifest filename format error",
            details_text=(
                f"Manifest file = {manifest_filename}; required format: {{BATCH}}.xlsx or YYYYMMDD_Danh sách mẫu_{{BATCH}}.xlsx"
            ),
        )

    # Manifest formatting checks: do not silently strip; fail if dirty.
    for _, r in manifest.iterrows():
        lid = r["LID_raw"]
        batch = r["Batch_raw"]
        rownum = r["Row_Number"]
        if lid == "":
            add_issue("Blank LID in manifest", row_number=rownum)
        elif lid != lid.strip():
            add_issue(
                "LID leading/trailing spaces",
                lid=lid,
                details_text=f"Raw LID value: {lid!r}",
                row_number=rownum,
            )
        if batch == "":
            add_issue("Blank batch code in manifest", lid=lid, row_number=rownum)
        elif batch != batch.strip():
            add_issue(
                "Batch code leading/trailing spaces",
                lid=lid,
                details_text=f"Raw batch value: {batch!r}",
                row_number=rownum,
            )

    manifest_clean = manifest.copy()
    manifest_clean["LID"] = manifest_clean["LID_raw"]
    manifest_clean["Batch"] = manifest_clean["Batch_raw"]

    # Batch code consistency.
    filename_batch = _batch_from_manifest_filename(manifest_path)
    batch_values = sorted([x for x in manifest_clean["Batch"].dropna().unique().tolist() if x != ""])
    if len(batch_values) != 1:
        add_issue("Manifest batch code not unique", details_text=f"Batch values found: {batch_values}")
        manifest_batch = batch_values[0] if batch_values else ""
    else:
        manifest_batch = batch_values[0]

    if filename_batch == "":
        add_issue("Cannot parse batch code from manifest filename", details_text=os.path.basename(manifest_path))
    elif manifest_batch and filename_batch != manifest_batch:
        # For metadata dir manifests the filename may include a date prefix
        # not present in MẺ CHẠY (e.g. 20260504_mtDNA_316 vs mtDNA_316).
        # Only flag as ERROR if the batch codes are genuinely unrelated.
        # Normalize spaces to underscores so that "mtDNA 321" matches "mtDNA_321".
        fn_norm = filename_batch.replace(" ", "_")
        mf_norm = manifest_batch.replace(" ", "_")
        if not (fn_norm.endswith(mf_norm) or mf_norm.endswith(fn_norm)):
            add_issue(
                "Batch code mismatch",
                details_text=f"Filename batch = {filename_batch}; Manifest batch = {manifest_batch}",
            )

    # Accept MS custom report matching either manifest_batch or filename_batch (metadata dir may use full BATCH_ID as filename).
    expected_batch_for_ms = manifest_batch or filename_batch
    ms_filename_ok = _custom_report_filename_matches_batch(ms_path, expected_batch_for_ms)
    if not ms_filename_ok and filename_batch and filename_batch != expected_batch_for_ms:
        ms_filename_ok = _custom_report_filename_matches_batch(ms_path, filename_batch)
        if ms_filename_ok:
            expected_batch_for_ms = filename_batch
    if not ms_filename_ok:
        add_issue(
            "Custom report filename format/batch error",
            details_text=(
                f"Custom report file = {ms_filename}; expected batch = {expected_batch_for_ms}; "
                "required format: {BATCH}_anything_df_custom_report.xlsx"
            ),
        )

    # Duplicate LID in manifest.
    nonblank_lids = [x for x in manifest_clean["LID"].tolist() if x != ""]
    dup_lids = (
        sorted(pd.Series(nonblank_lids)[pd.Series(nonblank_lids).duplicated()].unique().tolist())
        if nonblank_lids
        else []
    )
    for lid in dup_lids:
        rows = manifest_clean.loc[manifest_clean["LID"] == lid, "Row_Number"].tolist()
        add_issue("Duplicate LID in manifest", lid=lid, details_text=f"Rows: {rows}")

    manifest_lids = sorted(set(nonblank_lids))
    expected_n = len(manifest_lids)

    # MS-side required derived fields.
    ms = ms_df.copy()
    ms["LID"] = ms["LID"].apply(_as_display_string)
    ms["Primer"] = ms["Primer"].apply(_as_display_string).str.upper()

    blank_lid_rows = ms.index[ms["LID"] == ""].tolist()
    for idx in blank_lid_rows:
        add_issue(
            "MS row has blank/unparsed LID",
            details_text=f"DataFrame row: {idx + 2}; Sample Name: {ms.loc[idx, 'Sample Name']}",
        )

    blank_primer_rows = ms.index[ms["Primer"] == ""].tolist()
    for idx in blank_primer_rows:
        add_issue(
            "MS row has blank/unparsed primer",
            lid=ms.loc[idx, "LID"],
            details_text=f"DataFrame row: {idx + 2}; Sample Name: {ms.loc[idx, 'Sample Name']}",
        )

    unexpected_primers = sorted(
        [p for p in ms["Primer"].dropna().unique().tolist() if p not in EXPECTED_PRIMERS and p != ""]
    )
    for primer in unexpected_primers:
        count = int((ms["Primer"] == primer).sum())
        add_issue("Unexpected primer in MS report", primer=primer, details_text=f"Observed count: {count}")

    # Separate real sample LIDs from controls. The LID manifest intentionally contains
    # only real samples, while MS custom reports can include PC/NTC controls.
    ms_nonblank = ms[ms["LID"] != ""].copy()
    control_mask = ms_nonblank["LID"].apply(is_control_lid)
    ms_controls = ms_nonblank[control_mask].copy()
    ms_samples = ms_nonblank[~control_mask].copy()

    ms_sample_lids = sorted(set(ms_samples["LID"].tolist()))
    ms_control_lids = sorted(set(ms_controls["LID"].tolist()))

    missing_lids = sorted(set(manifest_lids) - set(ms_sample_lids))
    extra_lids = sorted(set(ms_sample_lids) - set(manifest_lids))

    for lid in missing_lids:
        add_issue(
            "Missing LID in MS report",
            lid=lid,
            details_text="Sample LID exists in manifest but not in MS custom report",
        )
    for lid in extra_lids:
        add_issue(
            "Extra non-control LID in MS report",
            lid=lid,
            details_text="Non-control LID exists in MS custom report but not in manifest",
        )

    # Primer total counts must equal manifest LID count for real samples only.
    # PC/NTC controls are excluded here and handled by the optional Control QC module.
    primer_rows = []
    for primer in EXPECTED_PRIMERS:
        observed = int((ms_samples["Primer"] == primer).sum())
        primer_rows.append(
            {
                "Primer": primer,
                "Expected_Count": expected_n,
                "Observed_Count": observed,
                "Result": "PASS" if observed == expected_n else "FAIL",
            }
        )
        if observed != expected_n:
            add_issue(
                "Primer total count mismatch",
                primer=primer,
                details_text=f"Expected {expected_n} sample traces; observed {observed} sample traces. PC/NTC controls are excluded.",
            )

    # Per-sample-LID x primer matrix must be exactly 1 for every expected primer.
    if not ms_samples.empty:
        counts = ms_samples.groupby(["LID", "Primer"]).size().reset_index(name="Count")
    else:
        counts = pd.DataFrame(columns=["LID", "Primer", "Count"])

    count_map = {(r["LID"], r["Primer"]): int(r["Count"]) for _, r in counts.iterrows()}
    lid_primer_rows = []
    for lid in manifest_lids:
        for primer in EXPECTED_PRIMERS:
            c = count_map.get((lid, primer), 0)
            result = "PASS" if c == 1 else "FAIL"
            lid_primer_rows.append(
                {"LID": lid, "Primer": primer, "Expected_Count": 1, "Observed_Count": c, "Result": result}
            )
            if c == 0:
                add_issue("Missing primer trace", lid=lid, primer=primer, details_text="Expected exactly 1 trace")
            elif c > 1:
                sample_names = (
                    ms_samples.loc[(ms_samples["LID"] == lid) & (ms_samples["Primer"] == primer), "Sample Name"]
                    .astype(str)
                    .tolist()
                )
                add_issue(
                    "Duplicate primer trace",
                    lid=lid,
                    primer=primer,
                    details_text=f"Observed {c} traces: {' | '.join(sample_names)}",
                )

    details_df = pd.DataFrame(details, columns=["Severity", "Issue_Type", "LID", "Primer", "Details", "Manifest_Row"])
    failed = len(details_df) > 0

    summary_rows = [
        {"Check": "QC Layer 0 Status", "Value": "FAIL" if failed else "PASS"},
        {"Check": "MS custom report file", "Value": os.path.basename(ms_path)},
        {"Check": "Required MS filename format", "Value": "{BATCH}_anything_df_custom_report.xlsx"},
        {"Check": "Manifest file", "Value": os.path.basename(manifest_path)},
        {"Check": "Required manifest filename format", "Value": "{BATCH}.xlsx or YYYYMMDD_Danh sách mẫu_{BATCH}.xlsx"},
        {"Check": "Manifest batch from filename", "Value": filename_batch},
        {"Check": "Manifest batch from MẺ CHẠY", "Value": manifest_batch},
        {"Check": "Expected LID count", "Value": expected_n},
        {"Check": "Observed MS sample LID count", "Value": len(ms_sample_lids)},
        {"Check": "Observed control LID count", "Value": len(ms_control_lids)},
        {"Check": "Observed control LIDs", "Value": "; ".join(ms_control_lids)},
        {"Check": "Missing LID count", "Value": len(missing_lids)},
        {"Check": "Extra non-control LID count", "Value": len(extra_lids)},
        {"Check": "Issue count", "Value": len(details_df)},
    ]
    if not details_df.empty:
        issue_counts = details_df["Issue_Type"].value_counts().reset_index()
        issue_counts.columns = ["Issue_Type", "Count"]
    else:
        issue_counts = pd.DataFrame(columns=["Issue_Type", "Count"])

    summary_df = pd.DataFrame(summary_rows)
    primer_counts_df = pd.DataFrame(primer_rows)
    lid_primer_df = pd.DataFrame(lid_primer_rows)

    with pd.ExcelWriter(inv_out, engine="openpyxl") as writer:
        summary_df.to_excel(writer, sheet_name="Inventory_Summary", index=False)
        issue_counts.to_excel(writer, sheet_name="Issue_Counts", index=False)
        details_df.to_excel(writer, sheet_name="Inventory_Details", index=False)
        primer_counts_df.to_excel(writer, sheet_name="Primer_Counts", index=False)
        lid_primer_df.to_excel(writer, sheet_name="LID_Primer_Matrix", index=False)

    print("====================================================")
    print("QC LAYER 0 - INVENTORY CHECK")
    print("====================================================")
    print(f"MS File        : {os.path.basename(ms_path)}")
    print(f"Manifest File  : {os.path.basename(manifest_path)}")
    print(f"Manifest Batch : {manifest_batch}")
    print(f"Expected LIDs  : {expected_n}")
    print(f"Observed sample LIDs : {len(ms_sample_lids)}")
    print(f"Observed control LIDs: {len(ms_control_lids)}")
    print("Result         : " + ("FAIL" if failed else "PASS"))
    print(f"Report         : {inv_out}")

    if failed:
        print("Issues Found:")
        for _, r in issue_counts.iterrows():
            print(f"  - {r['Issue_Type']}: {r['Count']}")
        print()
        print("THIS BATCH IS NOT ELIGIBLE FOR AUTO-PASS.")
        print("Please review AB1 files, Mutation Surveyor project, custom report export, and LID manifest.")
        print("No MS_trace_processed output generated.")
        print("====================================================")
        raise SystemExit(1)

    print("Inventory QC passed. Proceeding to MS_trace_processed generation.")
    print("====================================================")
    return inv_out


# =========================
# VARIANT SUMMARY
# =========================
def build_variant_summary(df):
    variant_cols = [c for c in df.columns if re.fullmatch(r"Variant\d+", c)]
    variant_cols = sorted(variant_cols, key=lambda x: int(re.search(r"\d+", x).group()))

    def build_summary(row):
        vals = []
        for c in variant_cols:
            v = row.get(c)
            if pd.isna(v):
                continue
            s = str(v).strip()
            if not s or s.lower() == "n.a.":
                continue
            vals.append(s)
        return " ".join(vals)

    df["Variant Summary"] = df.apply(build_summary, axis=1)
    df["Variant Summary"] = (
        df["Variant Summary"]
        .astype(str)
        .str.replace(r"\bc\.", "", regex=True)
        .str.replace(",", "", regex=False)
        .str.replace(r"\s+", " ", regex=True)
        .str.strip()
    )
    df.loc[df["Variant Summary"].isin(["nan", "None"]), "Variant Summary"] = ""
    return df


# =========================
# VARIANT CONVERTED
# =========================
def convert_token(tok):
    tok = re.sub(r"\$\d+", "", str(tok))
    tok = tok.replace(",", "")
    tok = re.sub(r"\bc\.", "", tok).strip()

    m_plain = re.match(r"^(\d+(?:\.\d+)?)([A-Za-z_]+)$", tok)
    if m_plain:
        return f"{m_plain.group(1)}{m_plain.group(2)}"

    m_sub = re.match(r"^(\d+(?:\.\d+)?)([ACGT])>([ACGT]+)$", tok)
    if m_sub:
        pos, ref, alt = m_sub.groups()
        key = "".join(sorted(set(alt)))
        code = UNORDERED_MAP.get(key, alt)
        return f"{pos}{code}"

    return tok


def build_variant_converted(df):
    df["Variant Converted"] = df["Variant Summary"].apply(
        lambda s: " ".join(convert_token(t) for t in str(s).split()) if pd.notna(s) else ""
    )
    return df


# =========================
# AC REPEAT (HVIII 513-525) - DICTIONARY/CONSENSUS-BASED
# =========================
RCRS_513_525 = {
    513: "G",
    514: "C",
    515: "A",
    516: "C",
    517: "A",
    518: "C",
    519: "A",
    520: "C",
    521: "A",
    522: "C",
    523: "A",
    524: "C",
    525: "C",
}
LOCAL_POSITIONS_513_525 = list(range(513, 526))
REF_SEQ_513_525 = "".join(RCRS_513_525[p] for p in LOCAL_POSITIONS_513_525)
REF_AC_CORE_513_525 = {
    515: "A",
    516: "C",
    517: "A",
    518: "C",
    519: "A",
    520: "C",
    521: "A",
    522: "C",
    523: "A",
    524: "C",
    525: "C",
}


def parse_local_ac_token(token):
    tok = str(token).strip().replace("_HET", "_het")
    m = re.match(r"^(\d+)([A-Za-z_]+)$", tok)
    if not m:
        return None
    pos = int(m.group(1))
    allele = m.group(2)

    if pos < 513 or pos > 525:
        return None

    if allele == "DEL":
        return {"pos": pos, "kind": "DEL", "value": "DEL", "token": tok}

    if re.fullmatch(r"[ACGT]", allele):
        return {"pos": pos, "kind": "SNP", "value": allele, "token": tok}

    return None


def split_local_ac_tokens(tokens):
    local_supported = []
    local_unsupported = []
    nonlocal_tokens = []

    for tok in tokens:
        parsed = parse_local_ac_token(tok)
        if parsed is not None:
            local_supported.append(tok)
            continue

        m = re.match(r"^(\d+(?:\.\d+)?)([A-Za-z_]+)$", str(tok))
        if m:
            try:
                pos = float(m.group(1))
            except Exception:
                pos = None
            if pos is not None and 513 <= pos <= 525:
                local_unsupported.append(tok)
                continue

        nonlocal_tokens.append(tok)

    return local_supported, local_unsupported, nonlocal_tokens


def build_local_ac_consensus(local_tokens):
    state = [RCRS_513_525[p] for p in LOCAL_POSITIONS_513_525]

    for tok in local_tokens:
        v = parse_local_ac_token(tok)
        if v is None:
            continue
        idx = v["pos"] - 513
        if v["kind"] == "DEL":
            state[idx] = None
        elif v["kind"] == "SNP":
            state[idx] = v["value"]

    return "".join(x for x in state if x is not None)


def enumerate_ac_alignments(ref_seq, target_seq):
    results = []

    def backtrack(i, j, path):
        if i == len(ref_seq) and j == len(target_seq):
            results.append(path.copy())
            return
        if i == len(ref_seq):
            return

        ref_base = ref_seq[i]
        ref_pos = 513 + i

        path.append(("DEL", ref_base, ref_pos, ""))
        backtrack(i + 1, j, path)
        path.pop()

        if j < len(target_seq):
            target_base = target_seq[j]
            if target_base == ref_base:
                path.append(("KEEP", ref_base, ref_pos, target_base))
            else:
                path.append(("SUB", ref_base, ref_pos, target_base))
            backtrack(i + 1, j + 1, path)
            path.pop()

    backtrack(0, 0, [])
    return results


def ac_motif_break_count(steps):
    breaks = 0
    for op, ref_base, pos, out_base in steps:
        if pos < 515 or pos > 525:
            continue
        if op == "DEL":
            continue
        expected = REF_AC_CORE_513_525[pos]
        if out_base != expected:
            breaks += 1
    return breaks


def deletion_positions(steps):
    return [pos for op, _, pos, _ in steps if op == "DEL"]


def substitution_positions(steps):
    return [pos for op, _, pos, _ in steps if op == "SUB"]


def contiguous_blocks(positions):
    if not positions:
        return 0
    positions = sorted(positions)
    blocks = 1
    for a, b in zip(positions[:-1], positions[1:]):
        if b != a + 1:
            blocks += 1
    return blocks


def score_ac_alignment(steps):
    dels = sorted(deletion_positions(steps))
    subs = sorted(substitution_positions(steps))
    right_align_key = tuple(-p for p in sorted(dels, reverse=True))

    return (
        ac_motif_break_count(steps),
        contiguous_blocks(dels),
        right_align_key,
        len(subs),
        tuple(-p for p in sorted(subs, reverse=True)),
    )


def ac_steps_to_tokens(steps):
    toks = []
    for op, ref_base, pos, out_base in steps:
        if op == "SUB":
            toks.append(f"{pos}{out_base}")
        elif op == "DEL":
            toks.append(f"{pos}DEL")
    return toks


def canonicalize_local_ac_consensus(consensus):
    alignments = enumerate_ac_alignments(REF_SEQ_513_525, consensus)
    if not alignments:
        return [], "no_valid_alignment"

    scored = sorted(alignments, key=score_ac_alignment)
    best = scored[0]
    best_tokens = ac_steps_to_tokens(best)

    return best_tokens, "dictionary_consensus_canonicalized"


def normalize_ac_repeat_tokens(tokens):
    """
    Dictionary/reference + local consensus + canonical mapping for 513-525.

    Supported local token forms:
    - SNP: 513A, 523G
    - DEL: 513DEL, 524DEL

    Unsupported local token forms inside 513-525:
    - insertions, e.g. 523.1C
    - heteroplasmy, e.g. 523het_C
    - other ambiguous forms

    Behavior:
    - if unsupported local tokens are present, leave AC-repeat region unchanged
    - otherwise canonicalize the local consensus and merge back with non-local tokens
    """
    toks = [str(t).strip() for t in tokens if str(t).strip()]
    local_supported, local_unsupported, nonlocal_tokens = split_local_ac_tokens(toks)

    if local_unsupported:
        return toks, []

    if not local_supported:
        return toks, []

    local_consensus = build_local_ac_consensus(local_supported)
    canonical_local_tokens, note = canonicalize_local_ac_consensus(local_consensus)

    merged = nonlocal_tokens + canonical_local_tokens

    def sort_key(tok):
        m = re.match(r"^(\d+(?:\.\d+)?)([A-Za-z_]+)$", str(tok))
        if not m:
            return (10**9, str(tok))
        return (float(m.group(1)), m.group(2))

    input_local_sorted = sorted(set(local_supported), key=sort_key)
    output_local_sorted = sorted(set(canonical_local_tokens), key=sort_key)
    merged_sorted = sorted(set(merged), key=sort_key)

    flags = [note] if input_local_sorted != output_local_sorted else []
    return merged_sorted, flags


# =========================
# VARIANT CONVERTED 2
# =========================
def build_variant_converted_2(df):
    """
    Build Variant Converted 2 using context-aware row-level normalization.

    C/T-stretch rules:
    - 303-309 -> 309 only if there is NO non-C variant anywhere in 303-309
    - 311-315 -> 315 only if there is NO non-C variant anywhere in 311-315
    - 452-455 -> 455 only if there is NO non-T variant anywhere in 452-455
      for insertion T tokens of the form xxx.yT
    - 568-573 -> 573 only if there is NO non-C variant anywhere in 568-573

    AC repeat rules:
    - Apply targeted HVIII 513-525 motif-preserving rewrites after the C/T-shift step.

    Debug columns:
    - C_shift_rule_applied
    - AC_repeat_rule_applied
    """

    def normalize_row(s):
        if pd.isna(s) or str(s).strip() == "":
            return ("", "", "")

        tokens = str(s).split()
        parsed = []

        for t in tokens:
            m = re.match(r"^(\d+(?:\.\d+)?)([A-Za-z_]+)$", t)
            if m:
                pos = float(m.group(1))
                allele = m.group(2)
                parsed.append((t, pos, allele))
            else:
                parsed.append((t, None, None))

        has_nonC_303_309 = any((pos is not None and 303 <= pos <= 309 and allele != "C") for _, pos, allele in parsed)

        has_nonC_311_315 = any((pos is not None and 311 <= pos <= 315 and allele != "C") for _, pos, allele in parsed)

        has_nonT_452_455 = any((pos is not None and 452 <= pos <= 455 and allele != "T") for _, pos, allele in parsed)

        has_nonC_568_573 = any((pos is not None and 568 <= pos <= 573 and allele != "C") for _, pos, allele in parsed)

        applied_303_309 = False
        applied_311_315 = False
        applied_452_455 = False
        applied_568_573 = False

        out = []
        for tok, pos, allele in parsed:
            m_c = re.match(r"^(\d+)(\.\d+)(C)$", tok)
            m_t = re.match(r"^(\d+)(\.\d+)(T)$", tok)

            if m_c:
                base = int(m_c.group(1))
                frac = m_c.group(2)

                if 303 <= base <= 309 and not has_nonC_303_309:
                    if base != 309:
                        applied_303_309 = True
                    base = 309
                elif 311 <= base <= 315 and not has_nonC_311_315:
                    if base != 315:
                        applied_311_315 = True
                    base = 315
                elif 568 <= base <= 573 and not has_nonC_568_573:
                    if base != 573:
                        applied_568_573 = True
                    base = 573

                out.append(f"{base}{frac}C")
                continue

            if m_t:
                base = int(m_t.group(1))
                frac = m_t.group(2)

                if 452 <= base <= 455 and not has_nonT_452_455:
                    if base != 455:
                        applied_452_455 = True
                    base = 455

                out.append(f"{base}{frac}T")
                continue

            out.append(tok)

        c_flags = []
        if applied_303_309:
            c_flags.append("303-309_to_309")
        if applied_311_315:
            c_flags.append("311-315_to_315")
        if applied_452_455:
            c_flags.append("452-455_to_455T")
        if applied_568_573:
            c_flags.append("568-573_to_573")

        if "248DEL" in out and "249DEL" not in out:
            out = ["249DEL" if tok == "248DEL" else tok for tok in out]
            c_flags.append("248DEL_to_249DEL")

        out, ac_flags = normalize_ac_repeat_tokens(out)

        return (" ".join(out), ";".join(c_flags), ";".join(ac_flags))

    results = df["Variant Converted"].apply(normalize_row)
    df["Variant Converted 2"] = results.apply(lambda x: x[0])
    df["C_shift_rule_applied"] = results.apply(lambda x: x[1])
    df["AC_repeat_rule_applied"] = results.apply(lambda x: x[2])
    return df


# =========================
# DERIVED COLUMNS
# =========================
def add_derived_columns(df):
    df["Well"] = df["Sample Name"].astype(str).str[:3]
    df["LID"] = df["Sample Name"].astype(str).str.extract(r"^[^_]+_\d{8}_(.*?)(?=_HV)", expand=False).fillna("")
    df["Primer"] = (
        df["Sample Name"]
        .astype(str)
        .str.extract(r"(HV1F|HV1R|HV2F|HV2R|HV3F|HV3R)", flags=re.IGNORECASE, expand=False)
        .fillna("")
        .str.upper()
    )

    df["MS Range"] = (
        df["Read Start"].astype(str).str.replace(r"\.0$", "", regex=True)
        + "-"
        + df["Read End"].astype(str).str.replace(r"\.0$", "", regex=True)
    )
    return df


# =========================
# QC COLUMNS
# =========================
def polyc_flag(row):
    """
    HV1 poly-C warning around 16180-16193.
    """
    primer = str(row.get("Primer", "")).strip().upper()
    if not primer.startswith("HV1"):
        return ""

    s = str(row.get("Variant Converted", ""))
    hits = 0
    for t in s.split():
        m = re.search(r"(\d+(?:\.\d+)?)", t)
        if not m:
            continue
        pos = float(m.group(1))
        if 16180 <= pos <= 16193:
            hits += 1

    if hits >= 5:
        return "BAD"
    if hits >= 2:
        return "SUSPECT"
    if hits >= 1:
        return "LOW"
    return ""


def c_stretch_flag(row):
    """
    HV2-3 C-stretch warning around 303-315.
    """
    primer = str(row.get("Primer", "")).strip().upper()
    if not (primer.startswith("HV2") or primer.startswith("HV3")):
        return ""

    s = str(row.get("Variant Converted", ""))
    hits = 0
    for t in s.split():
        m = re.match(r"^(\d+)\.\d+C$", t)
        if not m:
            continue
        pos = int(m.group(1))
        if 303 <= pos <= 315:
            hits += 1

    if hits >= 5:
        return "BAD"
    if hits >= 2:
        return "SUSPECT"
    if hits >= 1:
        return "LOW"
    return ""


# =========================
# FINAL COLUMNS
# =========================
def finalize_columns(df):
    keep_order = [
        "Well",
        "LID",
        "Primer",
        "Trace #",
        "Sample Name",
        "# of Mutation",
        "Read Start",
        "Read End",
        "MS Range",
        "Variant Summary",
        "Variant Converted",
        "Variant Converted 2",
        "C_shift_rule_applied",
        "AC_repeat_rule_applied",
        "HV1_PolyC_flag",
        "HV2-3_PolyC_flag",
    ]
    existing = [c for c in keep_order if c in df.columns]
    return df[existing]


# =========================
# MAIN
# =========================
def main():
    df = read_ms_all_blocks(MS_FILE)

    # Clean out non-real trace rows only.
    df = df[df["Sample Name"].astype(str).str.contains(".ab1", na=False)]
    df = df[df["Read Start"].notna() & df["Read End"].notna()]
    df = df.reset_index(drop=True)

    df = build_variant_summary(df)
    df = build_variant_converted(df)
    df = build_variant_converted_2(df)
    df = add_derived_columns(df)

    # Strict batch-level inventory gate. If this fails, the script exits before
    # creating MS_trace_processed output.
    run_inventory_qc(df, LID_MANIFEST, OUTPUT, MS_FILE)

    df["HV1_PolyC_flag"] = df.apply(polyc_flag, axis=1)
    df["HV2-3_PolyC_flag"] = df.apply(c_stretch_flag, axis=1)

    df = finalize_columns(df)
    df.to_excel(OUTPUT, index=False)

    print(f"Saved: {OUTPUT}")
    print(df["Primer"].value_counts(dropna=False))


if __name__ == "__main__":
    main()
