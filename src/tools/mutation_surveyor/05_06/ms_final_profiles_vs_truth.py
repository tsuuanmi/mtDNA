#!/usr/bin/env python3
"""
ms_final_profiles_vs_truth.py

Compare MS merged profiles against a truthset by normalized LID.

v5 adds region-level concordance checks on top of the v4 full-profile check:
    1) Full profile concordance:
       MS Sample profile - Variants
       vs Truth Sample Profile - Variants filtered to full region

    2) HV2-3 region concordance:
       MS Variants by HV2-3 Regions Adjusted
       vs Truth Sample Profile - Variants filtered to HV2-3 regions

    3) HV1 region concordance:
       MS Variants by HV1 Region Adjusted
       vs Truth Sample Profile - Variants filtered to HV1 region

Truth full region:
    73-340 438-576 16024-16365

HV2-3 truth region:
    73-340 438-576

HV1 truth region:
    16024-16365

Core concordance logic:
    - Exact canonical variant match => Y
    - Non-het IUPAC-compatible allele match => Y1
    - Het is treated as a distinct value
    - Het vs non-het => F
    - Equivalent het spellings are normalized before exact comparison
      e.g. 146C_het, 146_hetC -> 146_hetC
    - Y2 is not emitted
    - Missing/extra/incompatible variant => F
    - In HV2-3/full-profile comparisons, ignore length variants 309.xC and 573.xC
    - In HV1 comparison, do not apply HV1 polyC equivalence normalization

Output:
    Original MS sheet rows plus appended columns:
        Full Profile Concordance
        Full Profile Discordance reason
        HV2-3 Concordance
        HV2-3 Discordance reason
        HV1 Concordance
        HV1 Discordance reason
        Discordance Category
        Range (Truth)
        Variants (Truth)
        Flags (Truth)
        Range by HV2-3 Region (Truth)
        Variants by HV2-3 Region (Truth)
        Range by HV1 Region (Truth)
        Variants by HV1 Region (Truth)

Usage:
    python3 ms_final_profiles_vs_truth.py \
      --merged "All_37_final_profiles.xlsx" \
      --truth "All_37_Truthset.xlsx" \
      --output "All_37_final_profiles_vs_truth.xlsx"
"""

from __future__ import annotations

import argparse
import re
from collections.abc import Iterable
from pathlib import Path

import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

FULL_REGION_TEXT = "73-340 438-576 16024-16365"
HV23_REGION_TEXT = "73-340 438-576"
HV1_REGION_TEXT = "16024-16365"

FULL_PROFILE_REGIONS: list[tuple[float, float]] = [(73, 340), (438, 576), (16024, 16365)]
HV23_REGIONS: list[tuple[float, float]] = [(73, 340), (438, 576)]
HV1_REGIONS: list[tuple[float, float]] = [(16024, 16365)]

IUPAC = {
    "A": {"A"},
    "C": {"C"},
    "G": {"G"},
    "T": {"T"},
    "R": {"A", "G"},
    "Y": {"C", "T"},
    "S": {"G", "C"},
    "W": {"A", "T"},
    "K": {"G", "T"},
    "M": {"A", "C"},
    "B": {"C", "G", "T"},
    "D": {"A", "G", "T"},
    "H": {"A", "C", "T"},
    "V": {"A", "C", "G"},
    "N": {"A", "C", "G", "T"},
}

EMPTY_TOKENS = {"", "OK", "NAN", "NONE", "NULL", "-"}
GENERATED_COLUMNS = [
    "Concordance",
    "Discordance reason",  # column names from v4 outputs
    "Full Profile Concordance",
    "Full Profile Discordance reason",
    "HV2-3 Concordance",
    "HV2-3 Discordance reason",
    "HV1 Concordance",
    "HV1 Discordance reason",
    "Discordance Category",
    "Range (Truth)",
    "Variants (Truth)",
    "Flags (Truth)",
    "Range by HV2-3 Region (Truth)",
    "Variants by HV2-3 Region (Truth)",
    "Range by HV1 Region (Truth)",
    "Variants by HV1 Region (Truth)",
]


def norm_col_name(x: object) -> str:
    return re.sub(r"\s+", " ", str(x).strip().lower())


def find_column(df: pd.DataFrame, candidates: list[str], required: bool = True) -> str | None:
    norm_map = {norm_col_name(c): c for c in df.columns}
    for cand in candidates:
        key = norm_col_name(cand)
        if key in norm_map:
            return norm_map[key]
    for cand in candidates:
        parts = [p for p in re.split(r"[^a-zA-Z0-9]+", cand.lower()) if p]
        for c in df.columns:
            lc = norm_col_name(c)
            if all(p in lc for p in parts):
                return c
    if required:
        raise ValueError(
            "Cannot find required column. Candidates: "
            + ", ".join(candidates)
            + "\nAvailable columns: "
            + ", ".join(map(str, df.columns))
        )
    return None


def pick_sheet(xlsx_path: Path, preferred: str = "LID_no_PC_NTC") -> str:
    xls = pd.ExcelFile(xlsx_path)
    if preferred in xls.sheet_names:
        return preferred
    for s in xls.sheet_names:
        sl = s.lower()
        if "lid" in sl and "pc" in sl and "ntc" in sl:
            return s
    return xls.sheet_names[0]


def to_text(x: object) -> str:
    if pd.isna(x):
        return ""
    s = str(x).strip()
    if s.upper() in EMPTY_TOKENS:
        return ""
    return s


def normalize_lid(x: object) -> str:
    if pd.isna(x):
        return ""
    s = str(x).strip()
    if s == "":
        return ""
    if re.fullmatch(r"\d+\.0", s):
        s = s[:-2]
    return re.sub(r"\s+", "", s)


def expand_truth_range(x: object) -> str:
    s = to_text(x)
    if s.upper() == "FULL REGION":
        return FULL_REGION_TEXT
    return s


def variant_position(v: str) -> float | None:
    m = re.match(r"^\s*(\d+(?:\.\d+)?)", str(v))
    return float(m.group(1)) if m else None


def variant_key(v: str) -> str:
    m = re.match(r"^\s*(\d+(?:\.\d+)?)", str(v))
    return m.group(1) if m else str(v).strip()


def normalize_het_variant_token(v: str) -> str:
    s = str(v).strip()
    if not s:
        return ""
    s = re.sub(r"(?i)heteroplasmy", "het", s)
    s = re.sub(r"(?i)_?het_?", "_het", s)
    m = re.match(r"^(\d+(?:\.\d+)?)([A-Za-z]+)_het$", s, flags=re.IGNORECASE)
    if m:
        return f"{m.group(1)}_het{m.group(2).upper()}"
    m = re.match(r"^(\d+(?:\.\d+)?)_het([A-Za-z]+)$", s, flags=re.IGNORECASE)
    if m:
        return f"{m.group(1)}_het{m.group(2).upper()}"
    m = re.match(r"^(\d+(?:\.\d+)?)het_?([A-Za-z]+)$", str(v).strip(), flags=re.IGNORECASE)
    if m:
        return f"{m.group(1)}_het{m.group(2).upper()}"
    return str(v).strip().upper()


def is_het_variant(v: str) -> bool:
    return "_het" in normalize_het_variant_token(v).lower()


def variant_allele(v: str) -> str:
    v2 = normalize_het_variant_token(v)
    if "_het" in v2.lower():
        m = re.match(r"^\s*\d+(?:\.\d+)?_het(.+?)\s*$", v2, flags=re.IGNORECASE)
        return m.group(1).strip().upper() if m else ""
    m = re.match(r"^\s*\d+(?:\.\d+)?(.+?)\s*$", v2)
    return m.group(1).strip().upper() if m else ""


def parse_variants(profile: object) -> list[str]:
    s = to_text(profile)
    if not s:
        return []
    s = s.replace(";", " ").replace(",", " ")
    toks = [t.strip() for t in s.split() if t.strip()]
    return [t for t in toks if t.upper() not in EMPTY_TOKENS]


def in_regions(pos: float | None, regions: Iterable[tuple[float, float]]) -> bool:
    if pos is None:
        return False
    return any(start <= pos <= end for start, end in regions)


def filter_region_variants(profile: object, regions: Iterable[tuple[float, float]]) -> list[str]:
    return [v for v in parse_variants(profile) if in_regions(variant_position(v), regions)]


def is_ignored_hv23_length_variant(v: str) -> bool:
    v2 = str(v).strip().replace("_het", "")
    m = re.match(r"^(\d+\.\d+)([A-Za-z]+)$", v2)
    if not m:
        return False
    pos = float(m.group(1))
    allele = m.group(2).upper()
    return allele == "C" and ((309 < pos < 310) or (573 < pos < 574))


def comparable_variants(profile: object, regions: Iterable[tuple[float, float]], ignore_hv23_length: bool) -> list[str]:
    variants = [normalize_het_variant_token(v) for v in filter_region_variants(profile, regions)]
    if ignore_hv23_length:
        variants = [v for v in variants if not is_ignored_hv23_length_variant(v)]
    return variants


def compatibility_label(ms_v: str, truth_v: str) -> str | None:
    ms_norm = normalize_het_variant_token(ms_v)
    tr_norm = normalize_het_variant_token(truth_v)
    if ms_norm == tr_norm:
        return "Y"
    if is_het_variant(ms_norm) or is_het_variant(tr_norm):
        return None
    ms_allele = variant_allele(ms_norm)
    tr_allele = variant_allele(tr_norm)
    if ms_allele and tr_allele and ms_allele == tr_allele:
        return "Y"
    if ms_allele in IUPAC and tr_allele in IUPAC and (IUPAC[ms_allele] & IUPAC[tr_allele]):
        return "Y1"
    return None


def compare_profiles(
    ms_profile: object,
    truth_profile: object,
    regions: Iterable[tuple[float, float]],
    ignore_hv23_length: bool,
) -> tuple[str, str, list[dict]]:
    ms_vars = comparable_variants(ms_profile, regions, ignore_hv23_length)
    truth_vars = comparable_variants(truth_profile, regions, ignore_hv23_length)
    if not ms_vars and not truth_vars:
        return "Y", "", []

    ms_by_key: dict[str, list[str]] = {}
    truth_by_key: dict[str, list[str]] = {}
    for v in ms_vars:
        ms_by_key.setdefault(variant_key(v), []).append(v)
    for v in truth_vars:
        truth_by_key.setdefault(variant_key(v), []).append(v)

    def sort_key(k: str) -> float:
        try:
            return float(k)
        except ValueError:
            return 999999.0

    labels: list[str] = []
    reasons: list[str] = []
    details: list[dict] = []

    for key in sorted(set(ms_by_key) | set(truth_by_key), key=sort_key):
        ms_list = ms_by_key.get(key, [])
        tr_list = truth_by_key.get(key, [])

        if len(ms_list) == 1 and len(tr_list) == 1:
            lab = compatibility_label(ms_list[0], tr_list[0])
            if lab is None:
                labels.append("F")
                reason = f"Incompatible allele at {key}: MS={ms_list[0]}; Truth={tr_list[0]}"
                reasons.append(reason)
                details.append(
                    {
                        "Position": key,
                        "MS variant": ms_list[0],
                        "Truth variant": tr_list[0],
                        "Issue": "Incompatible allele",
                    }
                )
            else:
                labels.append(lab)
            continue

        if not ms_list and tr_list:
            labels.append("F")
            reason = f"Missing MS variant at {key}: Truth={' '.join(tr_list)}"
            reasons.append(reason)
            details.append(
                {"Position": key, "MS variant": "", "Truth variant": " ".join(tr_list), "Issue": "Missing MS variant"}
            )
            continue

        if ms_list and not tr_list:
            labels.append("F")
            reason = f"Extra MS variant at {key}: MS={' '.join(ms_list)}"
            reasons.append(reason)
            details.append(
                {"Position": key, "MS variant": " ".join(ms_list), "Truth variant": "", "Issue": "Extra MS variant"}
            )
            continue

        if sorted(ms_list) == sorted(tr_list):
            labels.append("Y")
            continue

        if len(ms_list) == len(tr_list):
            unmatched_truth = tr_list.copy()
            pair_labels: list[str] = []
            ok = True
            for ms_v in ms_list:
                matched_idx = None
                matched_label = None
                for i, tr_v in enumerate(unmatched_truth):
                    lab = compatibility_label(ms_v, tr_v)
                    if lab is not None:
                        matched_idx = i
                        matched_label = lab
                        break
                if matched_idx is None:
                    ok = False
                    break
                pair_labels.append(matched_label or "Y")
                unmatched_truth.pop(matched_idx)
            if ok and not unmatched_truth:
                labels.append("Y1" if "Y1" in pair_labels else "Y")
                continue

        labels.append("F")
        reason = f"Variant count/content mismatch at {key}: MS={' '.join(ms_list)}; Truth={' '.join(tr_list)}"
        reasons.append(reason)
        details.append(
            {
                "Position": key,
                "MS variant": " ".join(ms_list),
                "Truth variant": " ".join(tr_list),
                "Issue": "Variant count/content mismatch",
            }
        )

    if "F" in labels:
        final = "F"
    elif "Y1" in labels:
        final = "Y1"
    else:
        final = "Y"
    return final, "; ".join(reasons), details


def classify_discordance(hv23: str, hv1: str, full: str) -> str:
    hv23_bad = hv23 == "F"
    hv1_bad = hv1 == "F"
    full_bad = full == "F"
    if hv23_bad and hv1_bad:
        return "Both_HV23_and_HV1"
    if hv23_bad:
        return "HV23_only"
    if hv1_bad:
        return "HV1_only"
    if full_bad:
        return "Full_profile_only"
    return "Concordant"


def load_inputs(merged_file: Path, truth_file: Path) -> tuple[pd.DataFrame, pd.DataFrame, str]:
    merged_sheet = pick_sheet(merged_file, preferred="LID_no_PC_NTC")
    merged_df = pd.read_excel(merged_file, sheet_name=merged_sheet, dtype={"LID": object})
    truth_df = pd.read_excel(truth_file, dtype={"LID": object})
    return merged_df, truth_df, merged_sheet


def build_truth_table(truth_df: pd.DataFrame) -> tuple[pd.DataFrame, pd.DataFrame]:
    truth_lid_col = find_column(truth_df, ["LID"])
    truth_range_col = find_column(
        truth_df,
        ["Sample profile - Range", "Sample Profile - Range", "Sample Profiles - Range", "Range (Truth)"],
    )
    truth_var_col = find_column(
        truth_df,
        ["Sample profile - Variants", "Sample Profile - Variants", "Variants (Truth)"],
    )
    truth_flag_col = find_column(
        truth_df,
        ["Sample profile - Flag", "Sample Profile - Flag", "Flags", "Flag", "Flags (Truth)"],
        required=False,
    )

    truth_small = truth_df[[c for c in [truth_lid_col, truth_range_col, truth_var_col, truth_flag_col] if c]].copy()
    truth_small["__join_LID__"] = truth_small[truth_lid_col].map(normalize_lid)
    truth_small["Range (Truth)"] = truth_small[truth_range_col].map(expand_truth_range)
    truth_small["Variants (Truth)"] = truth_small[truth_var_col].map(to_text)
    truth_small["Flags (Truth)"] = truth_small[truth_flag_col].map(to_text) if truth_flag_col else ""
    truth_small["Range by HV2-3 Region (Truth)"] = HV23_REGION_TEXT
    truth_small["Variants by HV2-3 Region (Truth)"] = truth_small["Variants (Truth)"].map(
        lambda x: " ".join(filter_region_variants(x, HV23_REGIONS))
    )
    truth_small["Range by HV1 Region (Truth)"] = HV1_REGION_TEXT
    truth_small["Variants by HV1 Region (Truth)"] = truth_small["Variants (Truth)"].map(
        lambda x: " ".join(filter_region_variants(x, HV1_REGIONS))
    )

    keep = [
        "__join_LID__",
        "Range (Truth)",
        "Variants (Truth)",
        "Flags (Truth)",
        "Range by HV2-3 Region (Truth)",
        "Variants by HV2-3 Region (Truth)",
        "Range by HV1 Region (Truth)",
        "Variants by HV1 Region (Truth)",
    ]
    truth_small = truth_small[keep]
    truth_dup = truth_small[truth_small["__join_LID__"].duplicated(keep=False)].copy()
    truth_small = truth_small.drop_duplicates(subset=["__join_LID__"], keep="first")
    return truth_small, truth_dup


def build_output(merged_df: pd.DataFrame, truth_df: pd.DataFrame) -> tuple[pd.DataFrame, pd.DataFrame, pd.DataFrame]:
    merged_lid_col = find_column(merged_df, ["LID"])
    ms_full_range_col = find_column(
        merged_df,
        ["Sample profile - Range", "Sample Profile - Range", "Sample Profiles - Range"],
    )
    ms_full_var_col = find_column(merged_df, ["Sample profile - Variants", "Sample Profile - Variants"])
    ms_hv23_var_col = find_column(
        merged_df,
        ["Variants by HV2-3 Regions Adjusted", "Variants by HV2-3 Region Adjusted"],
    )
    ms_hv1_var_col = find_column(merged_df, ["Variants by HV1 Region Adjusted", "Variants by HV1 Regions Adjusted"])

    truth_small, truth_dup = build_truth_table(truth_df)

    out = merged_df.copy()
    out = out.drop(columns=[c for c in GENERATED_COLUMNS if c in out.columns], errors="ignore")
    out["__join_LID__"] = out[merged_lid_col].map(normalize_lid)
    out = out.merge(truth_small, how="left", on="__join_LID__")
    out = out.drop(columns=["__join_LID__"], errors="ignore")

    full_conc: list[str] = []
    full_reason: list[str] = []
    hv23_conc: list[str] = []
    hv23_reason: list[str] = []
    hv1_conc: list[str] = []
    hv1_reason: list[str] = []
    categories: list[str] = []
    detail_records: list[dict] = []

    for _, row in out.iterrows():
        lid = row.get(merged_lid_col, "")
        truth_profile = row.get("Variants (Truth)", "")
        truth_range = row.get("Range (Truth)", "")
        truth_flags = row.get("Flags (Truth)", "")

        if to_text(truth_profile) == "" and to_text(truth_range) == "":
            f_conc = h23_conc = h1_conc = "F"
            f_reason = h23_reason = h1_reason = "Missing truth row by LID"
            detail_sets = {
                "Full": [
                    {
                        "Position": "",
                        "MS variant": to_text(row.get(ms_full_var_col, "")),
                        "Truth variant": "",
                        "Issue": "Missing truth row by LID",
                    }
                ],
                "HV2-3": [
                    {
                        "Position": "",
                        "MS variant": to_text(row.get(ms_hv23_var_col, "")),
                        "Truth variant": "",
                        "Issue": "Missing truth row by LID",
                    }
                ],
                "HV1": [
                    {
                        "Position": "",
                        "MS variant": to_text(row.get(ms_hv1_var_col, "")),
                        "Truth variant": "",
                        "Issue": "Missing truth row by LID",
                    }
                ],
            }
        else:
            f_conc, f_reason, f_details = compare_profiles(
                row.get(ms_full_var_col, ""),
                truth_profile,
                FULL_PROFILE_REGIONS,
                True,
            )
            h23_conc, h23_reason, h23_details = compare_profiles(
                row.get(ms_hv23_var_col, ""),
                truth_profile,
                HV23_REGIONS,
                True,
            )
            h1_conc, h1_reason, h1_details = compare_profiles(
                row.get(ms_hv1_var_col, ""),
                truth_profile,
                HV1_REGIONS,
                False,
            )
            detail_sets = {"Full": f_details, "HV2-3": h23_details, "HV1": h1_details}

        full_conc.append(f_conc)
        full_reason.append(f_reason)
        hv23_conc.append(h23_conc)
        hv23_reason.append(h23_reason)
        hv1_conc.append(h1_conc)
        hv1_reason.append(h1_reason)
        categories.append(classify_discordance(h23_conc, h1_conc, f_conc))

        for scope, details in detail_sets.items():
            for d in details:
                detail_records.append(
                    {
                        "LID": lid,
                        "Scope": scope,
                        "Concordance": {"Full": f_conc, "HV2-3": h23_conc, "HV1": h1_conc}[scope],
                        "Discordance reason": {"Full": f_reason, "HV2-3": h23_reason, "HV1": h1_reason}[scope],
                        **d,
                        "MS Sample profile - Range": row.get(ms_full_range_col, ""),
                        "MS Sample profile - Variants": row.get(ms_full_var_col, ""),
                        "MS HV2-3 Variants Adjusted": row.get(ms_hv23_var_col, ""),
                        "MS HV1 Variants Adjusted": row.get(ms_hv1_var_col, ""),
                        "Truth profile - Range": truth_range,
                        "Truth profile - Variants": truth_profile,
                        "Truth HV2-3 Variants": row.get("Variants by HV2-3 Region (Truth)", ""),
                        "Truth HV1 Variants": row.get("Variants by HV1 Region (Truth)", ""),
                        "Truth flags": truth_flags,
                    }
                )

    out["Full Profile Concordance"] = full_conc
    out["Full Profile Discordance reason"] = full_reason
    out["HV2-3 Concordance"] = hv23_conc
    out["HV2-3 Discordance reason"] = hv23_reason
    out["HV1 Concordance"] = hv1_conc
    out["HV1 Discordance reason"] = hv1_reason
    out["Discordance Category"] = categories

    tail = [
        "Full Profile Concordance",
        "Full Profile Discordance reason",
        "HV2-3 Concordance",
        "HV2-3 Discordance reason",
        "HV1 Concordance",
        "HV1 Discordance reason",
        "Discordance Category",
        "Range (Truth)",
        "Variants (Truth)",
        "Flags (Truth)",
        "Range by HV2-3 Region (Truth)",
        "Variants by HV2-3 Region (Truth)",
        "Range by HV1 Region (Truth)",
        "Variants by HV1 Region (Truth)",
    ]
    base_cols = [c for c in out.columns if c not in tail]
    out = out[base_cols + tail]

    summary_rows = []
    for col in ["Full Profile Concordance", "HV2-3 Concordance", "HV1 Concordance", "Discordance Category"]:
        counts = out[col].fillna("").value_counts(dropna=False)
        for label, count in counts.items():
            summary_rows.append(
                {"Metric": col, "Label": label, "Count": int(count), "Percent": count / len(out) if len(out) else 0}
            )
    summary = pd.DataFrame(summary_rows)

    detail_df = pd.DataFrame(detail_records)
    if not truth_dup.empty:
        dup_records = []
        for _, r in truth_dup.iterrows():
            dup_records.append(
                {
                    "LID": r.get("__join_LID__", ""),
                    "Scope": "Truth",
                    "Concordance": "",
                    "Discordance reason": "Duplicate truth LID; first truth row used",
                    "Position": "",
                    "MS variant": "",
                    "Truth variant": r.get("Variants (Truth)", ""),
                    "Issue": "Duplicate truth LID; first truth row used",
                }
            )
        detail_df = pd.concat([detail_df, pd.DataFrame(dup_records)], ignore_index=True)
    return out, summary, detail_df


def autosize_and_style(xlsx_path: Path) -> None:
    wb = load_workbook(xlsx_path)
    header_fill = PatternFill("solid", fgColor="D9EAF7")
    bold = Font(bold=True)
    for ws in wb.worksheets:
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions
        for cell in ws[1]:
            cell.font = bold
            cell.fill = header_fill
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        for col_cells in ws.columns:
            max_len = 0
            col_letter = get_column_letter(col_cells[0].column)
            for cell in col_cells[:2000]:
                val = "" if cell.value is None else str(cell.value)
                max_len = max(max_len, len(val))
            ws.column_dimensions[col_letter].width = min(max(max_len + 2, 10), 60)
    wb.save(xlsx_path)


def main(merged_file: str, truth_file: str, output_file: str) -> None:
    merged_path = Path(merged_file)
    truth_path = Path(truth_file)
    output_path = Path(output_file)
    if not merged_path.exists():
        raise FileNotFoundError(f"Merged file not found: {merged_path}")
    if not truth_path.exists():
        raise FileNotFoundError(f"Truthset file not found: {truth_path}")
    merged_df, truth_df, merged_sheet = load_inputs(merged_path, truth_path)
    out_df, summary_df, detail_df = build_output(merged_df, truth_df)
    with pd.ExcelWriter(output_path, engine="openpyxl") as writer:
        out_df.to_excel(writer, sheet_name=merged_sheet[:31], index=False)
        summary_df.to_excel(writer, sheet_name="Summary", index=False)
        if detail_df.empty:
            detail_df = pd.DataFrame(
                columns=[
                    "LID",
                    "Scope",
                    "Concordance",
                    "Discordance reason",
                    "Position",
                    "MS variant",
                    "Truth variant",
                    "Issue",
                ]
            )
        detail_df.to_excel(writer, sheet_name="Discordance_detail", index=False)
    autosize_and_style(output_path)
    print("Done.")
    print(f"Input merged sheet: {merged_sheet}")
    print(f"Rows: {len(out_df)}")
    print("Summary:")
    print(summary_df.to_string(index=False))
    print(f"Output: {output_path}")


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(
        description="Compare final MS full and region-adjusted profiles against truthset by normalized LID."
    )
    parser.add_argument("--merged", required=True, help="Input *_final_profiles.xlsx from ms_final_profiles_merge.py")
    parser.add_argument("--truth", required=True, help="Truthset workbook")
    parser.add_argument("--output", required=True, help="Output *_final_profiles_vs_truth.xlsx workbook")
    return parser.parse_args()


if __name__ == "__main__":
    args = parse_args()
    main(args.merged, args.truth, args.output)
